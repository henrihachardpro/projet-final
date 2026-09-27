# -*- coding: utf-8 -*-
"""
SUFIX OptimiZer V1.1.3 - Moteur métier.

Fonctions principales :
- sélection d'un CSV SUFIX puis du PDF associé ;
- lecture du sommaire en page 2 ;
- extraction du nombre de supports sur la première page de chaque support ;
- confirmation/correction manuelle des nombres extraits ;
- création d'un nouveau classeur Excel ;
- nomenclature par support ;
- nomenclature totale rationalisée ;
- séparation longueurs / pièces ;
- choix des longueurs de barres disponibles ;
- optimisation des découpes ;
- synthèse globale.

Le fichier CSV d'origine n'est jamais modifié.
"""

from __future__ import annotations

import csv
import json
import math
import re
import sys
import time
import traceback
import unicodedata
from collections import defaultdict
from datetime import datetime
from difflib import SequenceMatcher
from functools import lru_cache
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional

import pdfplumber
try:
    import fitz  # PyMuPDF : moteur PDF rapide
except ImportError:  # secours si dépendance indisponible
    fitz = None
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import tkinter as tk
from tkinter import filedialog, messagebox, ttk


APP_TITLE = "Sufix OptimiZer"
PROFILE_DATABASE_FILE = "sufix_profiles.xlsx"
ARTICLE_DATABASE_FILE = "base_article_supportage.xlsx"
VERSION_FILE = "version.txt"
SPLASH_FILE = "splash_screen.png"
WELCOME_BACKGROUND_FILE = "welcome_background.png"
RELEASE_MONTH_LABEL = "Août 2026"
ICON_FILE = "app_icon.ico"
ERP_COMMENT_LINE_TYPE = "Comment"
SHEET_GLOBAL = "Etude Globale"
SHEET_BY_SUPPORT = "Nomenclature par support"
SHEET_TOTAL = "Nomenclature totale pour opti"
SHEET_LENGTHS = "Opti longueurs"
SHEET_PIECES = "Opti pcs"
SHEET_INITIAL = "Offre initiale"
SHEET_SUMMARY = "Opti globale"
SHEET_CUT_PLAN = "Plan de découpe"
SHEET_SETTINGS = "_Paramètres"
SHEET_CONTROLS = "Contrôles"
SHEET_CABLE_TRAY = "Opti Chemin de câbles"
SHEET_COMPARISON = "Synthèse optimisation"
DATA_VERSION_FILE = "data_versions.json"
CABLE_TRAY_STOCK_LENGTH_M = 3.0

COLUMN_ALIASES = {
    "niveau": "Niveau",
    "level": "Niveau",
    "support": "Support",
    "code article": "Code article",
    "codearticle": "Code article",
    "article": "Code article",
    "libelle": "Libellé",
    "libelle article": "Libellé",
    "designation": "Libellé",
    "quantite": "Quantité",
    "qte": "Quantité",
    "quantite totale": "Quantité",
    "qte totale": "Quantité",
    "longueur utile": "Longueur utile",
    "longueur proposee": "Longueur proposée",
    "chute calculee": "Chute calculée",
    "nombre de supports": "Nombre de supports",
    "nombre supports": "Nombre de supports",
    "nb supports": "Nombre de supports",
    "nb de supports": "Nombre de supports",
    "nb support": "Nombre de supports",
    "quantite pour 1 support": "Quantité pour 1 support",
    "quantite pour un support": "Quantité pour 1 support",
    "quantite par support": "Quantité pour 1 support",
    "qte par support": "Quantité pour 1 support",
    "qte pour 1 support": "Quantité pour 1 support",
    "quantite unitaire support": "Quantité pour 1 support",
    "espacement entre support": "Espacement entre supports (m)",
    "espacement entre supports": "Espacement entre supports (m)",
    "espacement entre les supports": "Espacement entre supports (m)",
    "espacement support": "Espacement entre supports (m)",
    "espacement supports": "Espacement entre supports (m)",
    "espace entre support": "Espacement entre supports (m)",
    "espace entre supports": "Espacement entre supports (m)",
    "espace entre les supports": "Espacement entre supports (m)",
}

REQUIRED_COLUMNS = [
    "Niveau",
    "Support",
    "Code article",
    "Libellé",
    "Quantité",
    "Longueur utile",
    "Longueur proposée",
    "Chute calculée",
]

OUTPUT_GLOBAL_COLUMNS = REQUIRED_COLUMNS + [
    "Nombre de supports",
    "Espacement entre supports (m)",
    "Quantité pour 1 support",
]

LENGTH_COLUMNS = ["Longueur utile", "Longueur proposée", "Chute calculée"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
SUBTITLE_FILL = PatternFill("solid", fgColor="DDEBF7")
ALT_ROW_FILL = PatternFill("solid", fgColor="EAF2F8")
WHITE_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=13, bold=True, color="1F1F1F")
THIN_GREY = Side(style="thin", color="B7B7B7")
TABLE_BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


@dataclass(frozen=True)
class SupportInfo:
    support: str
    page: Optional[int]
    count: int
    level: str = ""
    source: str = "PDF"
    spacing_m: Optional[float] = None
    spacing_source: str = ""
    chassis: bool = False


@dataclass(frozen=True)
class ProfileOption:
    profile_key: str
    profile_label: str
    category: str
    stock_lengths_m: tuple[float, ...]


@dataclass
class CutBin:
    stock_length_mm: int
    cuts_mm: list[int]

    @property
    def used_mm(self) -> int:
        return sum(self.cuts_mm)

    @property
    def waste_mm(self) -> int:
        return self.stock_length_mm - self.used_mm


def resource_path(relative_name: str) -> Path:
    """Retourne le chemin d'une ressource, y compris dans un exécutable PyInstaller."""
    if hasattr(sys, "_MEIPASS"):
        return Path(getattr(sys, "_MEIPASS")) / relative_name
    return Path(__file__).resolve().parent / relative_name


def read_application_version() -> str:
    version_path = resource_path(VERSION_FILE)
    try:
        content = version_path.read_text(encoding="utf-8").strip()
    except OSError:
        return "1.0.0"
    match = re.search(r"(?i)version\s*:\s*([0-9]+(?:\.[0-9]+){1,3})", content)
    return match.group(1) if match else content.replace("Version", "").replace(":", "").strip()



def read_data_versions() -> dict[str, object]:
    try:
        return json.loads(resource_path(DATA_VERSION_FILE).read_text(encoding="utf-8"))
    except Exception:
        return {}


def set_window_icon(window: tk.Misc) -> None:
    try:
        window.iconbitmap(default=str(resource_path(ICON_FILE)))
    except Exception:
        pass


@lru_cache(maxsize=16384)
def _normalize_text_cached(text: str) -> str:
    normalized = text.replace("\u00a0", " ").replace("–", "-").replace("—", "-")
    normalized = unicodedata.normalize("NFKD", normalized)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", normalized).strip().lower()


def normalize_text(value: object) -> str:
    return _normalize_text_cached("" if value is None else str(value))


def clean_cell(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if stripped in {"-", "–", "—"}:
            return None
        return stripped
    return value


def parse_float(value: object) -> Optional[float]:
    value = clean_cell(value)
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    text = text.replace(",", ".")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if not match:
        return None
    return float(match.group(0))


def display_number(value: object) -> object:
    """
    Met en forme uniquement les valeurs déjà numériques.

    Une chaîne telle que "CIN24208015", "Niveau 1" ou
    "Collier SX242 Ø 76,1" doit rester du texte.
    """
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if abs(value - round(value)) < 1e-9:
            return int(round(value))
        return round(value, 6)
    return value


def decode_csv_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    encodings = ["utf-8-sig", "cp1252", "latin-1"]
    for encoding in encodings:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError("Encodage du CSV non reconnu.")


def canonicalize_header(header: object) -> str:
    raw = "" if header is None else str(header).strip()
    normalized = normalize_text(raw)

    if normalized in COLUMN_ALIASES:
        return COLUMN_ALIASES[normalized]

    # Détection souple des deux colonnes susceptibles d'être ajoutées par SUFIX.
    # L'objectif est de tolérer des variantes de nom sans dépendre de leur position.
    if "support" in normalized:
        if (
            re.search(r"\b(nb|nombre)\b", normalized)
            and "quant" not in normalized
            and "qte" not in normalized
        ):
            return "Nombre de supports"

        quantity_words = ("quantite", "qte")
        per_support_words = ("par support", "pour 1 support", "pour un support", "unitaire")
        if any(word in normalized for word in quantity_words) and any(
            word in normalized for word in per_support_words
        ):
            return "Quantité pour 1 support"

    return raw


def canonicalize_headers(headers: Iterable[str]) -> list[str]:
    result: list[str] = []
    used: dict[str, int] = {}
    for header in headers:
        canonical = canonicalize_header(header)
        if canonical in used:
            used[canonical] += 1
            canonical = f"{canonical} ({used[canonical]})"
        else:
            used[canonical] = 1
        result.append(canonical)
    return result


def _parse_csv_candidate(
    lines: list[str],
    delimiter: str,
    max_header_rows: int = 15,
) -> tuple[int, list[str], list[list[str]], int]:
    parsed = list(csv.reader(lines, delimiter=delimiter))
    best_index = -1
    best_headers: list[str] = []
    best_score = -1

    for index, values in enumerate(parsed[:max_header_rows]):
        if not values or all(not str(value).strip() for value in values):
            continue
        headers = canonicalize_headers(values)
        score = sum(1 for column in REQUIRED_COLUMNS if column in headers)
        if "Nombre de supports" in headers:
            score += 1
        if "Quantité pour 1 support" in headers:
            score += 1
        if score > best_score:
            best_index = index
            best_headers = headers
            best_score = score

    return best_index, best_headers, parsed, best_score


def detect_csv_structure(path: Path) -> tuple[str, str, int, list[str], list[list[str]]]:
    text, encoding = decode_csv_text(path)
    lines = text.splitlines()
    if not lines:
        raise ValueError("Le fichier CSV est vide.")

    # Excel/SUFIX peut placer une directive spéciale en première ligne : sep=,
    # ou sep=; . Elle n'est pas une ligne d'en-tête.
    first_nonempty_index = next(
        (i for i, line in enumerate(lines) if line.strip()),
        None,
    )
    directive_delimiter: Optional[str] = None
    content_start = 0
    if first_nonempty_index is not None:
        match = re.match(r"^\s*sep\s*=\s*(.)\s*$", lines[first_nonempty_index], re.I)
        if match:
            directive_delimiter = match.group(1)
            content_start = first_nonempty_index + 1

    content_lines = lines[content_start:]
    delimiter_candidates = (
        [directive_delimiter]
        if directive_delimiter
        else [",", ";", "\t", "|"]
    )

    candidates = []
    for delimiter in delimiter_candidates:
        header_index, headers, parsed, score = _parse_csv_candidate(
            content_lines,
            delimiter,
        )
        candidates.append((score, delimiter, header_index, headers, parsed))

    # Si aucune directive sep= n'était fournie, essayer malgré tout tous les
    # séparateurs et choisir celui qui reconnaît le mieux le schéma SUFIX.
    if directive_delimiter is None:
        candidates.sort(key=lambda item: item[0], reverse=True)
    best_score, delimiter, header_index, headers, parsed = max(
        candidates,
        key=lambda item: item[0],
    )

    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        detected = ", ".join(header for header in headers if header) or "(aucune)"
        raise ValueError(
            "Impossible d'identifier correctement les colonnes SUFIX.\n\n"
            "Colonnes obligatoires manquantes : "
            + ", ".join(missing)
            + "\n\nColonnes détectées : "
            + detected
        )

    return encoding, delimiter, header_index, headers, parsed


def read_sufix_csv(path: Path) -> list[dict[str, object]]:
    _, _, header_index, headers, parsed = detect_csv_structure(path)
    rows: list[dict[str, object]] = []

    for values in parsed[header_index + 1 :]:
        if not values or all(not str(value).strip() for value in values):
            continue
        if len(values) < len(headers):
            values = list(values) + [""] * (len(headers) - len(values))

        row = {
            headers[index]: clean_cell(values[index])
            for index in range(min(len(headers), len(values)))
        }

        for text_column in ("Niveau", "Support", "Code article", "Libellé"):
            value = row.get(text_column)
            row[text_column] = "" if value is None else str(value)

        row["Quantité"] = parse_float(row.get("Quantité")) or 0.0
        for column in LENGTH_COLUMNS:
            row[column] = parse_float(row.get(column))

        if "Nombre de supports" in row:
            row["Nombre de supports"] = parse_float(row.get("Nombre de supports"))
        if "Quantité pour 1 support" in row:
            row["Quantité pour 1 support"] = parse_float(
                row.get("Quantité pour 1 support")
            )
        if "Espacement entre supports (m)" in row:
            row["Espacement entre supports (m)"] = parse_float(
                row.get("Espacement entre supports (m)")
            )

        rows.append(row)

    if not rows:
        raise ValueError("Le CSV ne contient aucune ligne de nomenclature.")
    return rows


def resolve_support_counts_from_csv(
    rows: list[dict[str, object]],
) -> tuple[dict[str, int], dict[str, str], list[str]]:
    """
    Tente de déterminer le nombre de supports sans PDF.

    Priorité :
    1. colonne explicite « Nombre de supports » ;
    2. déduction Quantité totale / Quantité pour 1 support.

    Retourne :
    - les nombres de supports fiables ;
    - la source de chaque valeur ;
    - la liste des supports restant à résoudre avec le PDF.
    """
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("Support") or "")].append(row)

    resolved: dict[str, int] = {}
    sources: dict[str, str] = {}
    unresolved: list[str] = []

    for support, support_rows in grouped.items():
        explicit_counts: list[int] = []
        explicit_invalid = False

        for row in support_rows:
            raw_count = parse_float(row.get("Nombre de supports"))
            if raw_count is None:
                continue
            rounded = round(raw_count)
            if raw_count <= 0 or abs(raw_count - rounded) > 1e-6:
                explicit_invalid = True
                break
            explicit_counts.append(int(rounded))

        if explicit_counts and not explicit_invalid:
            unique_counts = set(explicit_counts)
            if len(unique_counts) == 1:
                count = explicit_counts[0]

                # Si une quantité par support existe également, contrôler sa cohérence.
                consistent = True
                for row in support_rows:
                    quantity_per_support = parse_float(
                        row.get("Quantité pour 1 support")
                    )
                    if quantity_per_support is None:
                        continue
                    total_quantity = parse_float(row.get("Quantité")) or 0.0
                    if not math.isclose(
                        total_quantity,
                        quantity_per_support * count,
                        rel_tol=1e-7,
                        abs_tol=1e-6,
                    ):
                        consistent = False
                        break

                if consistent:
                    resolved[support] = count
                    sources[support] = "CSV - nombre de supports"
                    continue

        # Deuxième possibilité : calculer le nombre de supports à partir de
        # Quantité / Quantité pour 1 support. Toutes les lignes exploitables
        # d'un même support doivent aboutir au même entier.
        derived_counts: list[int] = []
        derivation_invalid = False
        for row in support_rows:
            total_quantity = parse_float(row.get("Quantité"))
            quantity_per_support = parse_float(row.get("Quantité pour 1 support"))
            if (
                total_quantity is None
                or quantity_per_support is None
                or quantity_per_support <= 0
            ):
                continue
            ratio = total_quantity / quantity_per_support
            rounded = round(ratio)
            if ratio <= 0 or abs(ratio - rounded) > 1e-6:
                derivation_invalid = True
                break
            derived_counts.append(int(rounded))

        if derived_counts and not derivation_invalid and len(set(derived_counts)) == 1:
            resolved[support] = derived_counts[0]
            sources[support] = "CSV - quantité par support"
        else:
            unresolved.append(support)

    return resolved, sources, sorted(unresolved, key=natural_number)



def resolve_support_spacings_from_csv(
    rows: list[dict[str, object]],
) -> tuple[dict[str, float], dict[str, str], list[str]]:
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("Support") or "")].append(row)
    resolved: dict[str, float] = {}
    sources: dict[str, str] = {}
    unresolved: list[str] = []
    for support, support_rows in grouped.items():
        values = [
            parse_float(row.get("Espacement entre supports (m)"))
            for row in support_rows
            if parse_float(row.get("Espacement entre supports (m)")) is not None
        ]
        values = [float(v) for v in values if v is not None and v > 0]
        rounded = {round(v, 6) for v in values}
        if rounded and len(rounded) == 1:
            resolved[support] = values[0]
            sources[support] = "CSV - espacement"
        else:
            unresolved.append(support)
    return resolved, sources, sorted(unresolved, key=natural_number)


def natural_number(text: object) -> int:
    match = re.search(r"(\d+)", str(text))
    return int(match.group(1)) if match else 10**9


def ordered_unique_supports(rows: list[dict[str, object]]) -> list[str]:
    """Conserve exactement l'ordre d'apparition des supports dans le CSV."""
    result: list[str] = []
    seen: set[str] = set()
    for row in rows:
        support = str(row.get("Support") or "").strip()
        key = normalize_text(support)
        if support and key not in seen:
            seen.add(key)
            result.append(support)
    return result


def _toc_entries_from_text(toc_text: str) -> list[tuple[int, str]]:
    """Extrait les couples (page, titre) génériques d'un sommaire SUFIX."""
    entries: list[tuple[int, str]] = []
    for raw_line in toc_text.replace("\u00a0", " ").splitlines():
        line = raw_line.strip()
        match = re.match(r"^(\d{1,4})\s+(.+?)\s*$", line)
        if not match:
            continue
        page = int(match.group(1))
        title = match.group(2).strip()
        if page > 0 and title:
            entries.append((page, title))
    return entries


def _support_name_on_first_page(text: str, support: str) -> bool:
    """
    Recherche un titre de support dans le début d'une page sans accepter les
    sous-chaînes ambiguës. Par exemple « CTA NC » ne doit pas correspondre à
    « UE BATTERIE CTA NC + C1 ».
    """
    target = normalize_text(support)
    lines = [
        normalize_text(line)
        for line in text.replace("\u00a0", " ").splitlines()
        if normalize_text(line)
    ]
    # La zone titre se trouve au début du texte extrait sur les pages support.
    head = lines[:18]
    for width in (1, 2, 3, 4):
        for index in range(0, max(0, len(head) - width + 1)):
            candidate = normalize_text(" ".join(head[index:index + width]))
            if candidate == target:
                return True
    return False


def _extract_support_pages_from_pdf(
    pdf,
    expected_supports: Optional[list[str]] = None,
    page_text_cache: Optional[dict[int, str]] = None,
) -> dict[str, int]:
    expected = [str(value).strip() for value in (expected_supports or []) if str(value).strip()]
    normalized_expected = {normalize_text(value): value for value in expected}
    cache = page_text_cache if page_text_cache is not None else {}
    if len(pdf.pages) < 2:
        raise ValueError("Le PDF doit contenir au moins deux pages.")
    toc_text = cache.get(2)
    if toc_text is None:
        toc_text = pdf.pages[1].extract_text() or ""
        cache[2] = toc_text
    entries = _toc_entries_from_text(toc_text)
    support_pages: dict[str, int] = {}
    if expected:
        for page, title in entries:
            support = normalized_expected.get(normalize_text(title))
            if support is not None:
                support_pages.setdefault(support, page)
        unresolved = [s for s in expected if s not in support_pages]
        for support in unresolved:
            target = normalize_text(support)
            candidates: list[tuple[float, int, str]] = []
            for page, title in entries:
                candidate = normalize_text(title)
                score = SequenceMatcher(None, target, candidate).ratio()
                if score >= 0.93:
                    candidates.append((score, page, title))
            if candidates:
                candidates.sort(key=lambda item: (-item[0], item[1]))
                best = candidates[0]
                if len(candidates) == 1 or best[0] - candidates[1][0] >= 0.03:
                    support_pages[support] = best[1]
        unresolved = [s for s in expected if s not in support_pages]
        if unresolved:
            for page_number, page_obj in enumerate(pdf.pages[2:], start=3):
                text = cache.get(page_number)
                if text is None:
                    text = page_obj.extract_text() or ""
                    cache[page_number] = text
                for support in list(unresolved):
                    if _support_name_on_first_page(text, support):
                        support_pages[support] = page_number
                        unresolved.remove(support)
                if not unresolved:
                    break
        return {support: support_pages[support] for support in expected if support in support_pages}
    for page, title in entries:
        if normalize_text(title).startswith("support "):
            support_pages[title.strip()] = page
    if not support_pages:
        for index, page_obj in enumerate(pdf.pages[2:], start=3):
            text = cache.get(index)
            if text is None:
                text = page_obj.extract_text() or ""
                cache[index] = text
            match = re.search(r"(?im)^\s*(SUPPORT\s+\d+)\s*$", text)
            if match:
                support_pages.setdefault(match.group(1).title(), index)
    return dict(sorted(support_pages.items(), key=lambda item: item[1]))


def _extract_support_pages_pdfplumber(pdf_path: Path, expected_supports: Optional[list[str]] = None) -> dict[str, int]:
    with pdfplumber.open(pdf_path) as pdf:
        return _extract_support_pages_from_pdf(pdf, expected_supports)


def _fitz_page_text(document, page_number: int, cache: dict[int, str]) -> str:
    text = cache.get(page_number)
    if text is None:
        text = document[page_number - 1].get_text("text") or ""
        cache[page_number] = text
    return text


def _fitz_words_as_pdfplumber(document, page_number: int) -> list[dict[str, object]]:
    words = []
    for item in document[page_number - 1].get_text("words") or []:
        # PyMuPDF : x0, y0, x1, y1, word, block, line, word_no
        x0, y0, x1, y1, text = item[:5]
        words.append({
            "text": text,
            "x0": x0,
            "x1": x1,
            "top": y0,
            "bottom": y1,
        })
    return words


def _extract_support_pages_fitz(
    document,
    expected_supports: Optional[list[str]] = None,
    page_text_cache: Optional[dict[int, str]] = None,
) -> dict[str, int]:
    expected = [str(value).strip() for value in (expected_supports or []) if str(value).strip()]
    normalized_expected = {normalize_text(value): value for value in expected}
    cache = page_text_cache if page_text_cache is not None else {}
    if document.page_count < 2:
        raise ValueError("Le PDF doit contenir au moins deux pages.")

    entries = _toc_entries_from_text(_fitz_page_text(document, 2, cache))
    support_pages: dict[str, int] = {}
    if expected:
        for page, title in entries:
            support = normalized_expected.get(normalize_text(title))
            if support is not None:
                support_pages.setdefault(support, page)

        unresolved = [s for s in expected if s not in support_pages]
        for support in unresolved:
            target = normalize_text(support)
            candidates: list[tuple[float, int, str]] = []
            for page, title in entries:
                candidate = normalize_text(title)
                score = SequenceMatcher(None, target, candidate).ratio()
                if score >= 0.93:
                    candidates.append((score, page, title))
            if candidates:
                candidates.sort(key=lambda item: (-item[0], item[1]))
                best = candidates[0]
                if len(candidates) == 1 or best[0] - candidates[1][0] >= 0.03:
                    support_pages[support] = best[1]

        unresolved = [s for s in expected if s not in support_pages]
        if unresolved:
            for page_number in range(3, document.page_count + 1):
                text = _fitz_page_text(document, page_number, cache)
                for support in list(unresolved):
                    if _support_name_on_first_page(text, support):
                        support_pages[support] = page_number
                        unresolved.remove(support)
                if not unresolved:
                    break
        return {support: support_pages[support] for support in expected if support in support_pages}

    for page, title in entries:
        if normalize_text(title).startswith("support "):
            support_pages[title.strip()] = page
    if not support_pages:
        for page_number in range(3, document.page_count + 1):
            text = _fitz_page_text(document, page_number, cache)
            match = re.search(r"(?im)^\s*(SUPPORT\s+\d+)\s*$", text)
            if match:
                support_pages.setdefault(match.group(1).title(), page_number)
    return dict(sorted(support_pages.items(), key=lambda item: item[1]))


def extract_support_pages(pdf_path: Path, expected_supports: Optional[list[str]] = None) -> dict[str, int]:
    if fitz is not None:
        try:
            with fitz.open(pdf_path) as document:
                return _extract_support_pages_fitz(document, expected_supports)
        except Exception:
            # Robustesse : le moteur historique reste disponible en secours.
            pass
    return _extract_support_pages_pdfplumber(pdf_path, expected_supports)


def extract_count_from_page(
    text: str,
    words: Optional[list[dict[str, object]]] = None,
) -> Optional[int]:
    """
    Extrait uniquement une valeur clairement rattachée au libellé
    « Nombre de supports ».

    Contrairement à l'ancienne recherche large, cette fonction ne cherche plus
    le premier nombre dans les dizaines de caractères suivantes. Cela évite de
    confondre le nombre de supports avec un diamètre, une longueur ou une quantité
    de nomenclature (par exemple 114 mm).
    """
    normalized_text = text.replace("\u00a0", " ")
    lines = [line.strip() for line in normalized_text.splitlines() if line.strip()]

    # Cas 1 : libellé et valeur sur la même ligne.
    for line in lines:
        match = re.fullmatch(
            r"(?i)\s*Nombre\s+de\s+supports\s*[:\-]?\s*(\d{1,5})\s*",
            line,
        )
        if match:
            value = int(match.group(1))
            return value if value > 0 else None

    # Cas 2 : la valeur est sur la ligne immédiatement située sous le libellé.
    for index, line in enumerate(lines):
        if normalize_text(line) != "nombre de supports":
            continue
        # Regarder au maximum les deux lignes suivantes, mais accepter uniquement
        # une ligne constituée d'un entier seul.
        for candidate in lines[index + 1 : index + 3]:
            match = re.fullmatch(r"\s*(\d{1,5})\s*", candidate)
            if match:
                value = int(match.group(1))
                return value if value > 0 else None
            # Dès qu'on rencontre un nouveau libellé textuel, arrêter : on ne doit
            # surtout pas aller chercher un nombre plus loin dans la nomenclature.
            if re.search(r"[A-Za-zÀ-ÿ]", candidate):
                break

    # Cas 3 : secours spatial grâce aux coordonnées de pdfplumber.
    # La valeur doit être un entier très proche verticalement du libellé.
    if words:
        ordered = sorted(
            words,
            key=lambda word: (
                float(word.get("top", 0) or 0),
                float(word.get("x0", 0) or 0),
            ),
        )
        for idx in range(len(ordered) - 2):
            trio = ordered[idx : idx + 3]
            trio_text = [normalize_text(word.get("text", "")) for word in trio]
            if trio_text != ["nombre", "de", "supports"]:
                continue

            tops = [float(word.get("top", 0) or 0) for word in trio]
            if max(tops) - min(tops) > 4:
                continue

            label_bottom = max(float(word.get("bottom", 0) or 0) for word in trio)
            label_x0 = min(float(word.get("x0", 0) or 0) for word in trio)
            label_x1 = max(float(word.get("x1", 0) or 0) for word in trio)

            candidates = []
            for word in ordered:
                raw = str(word.get("text", "")).strip()
                if not re.fullmatch(r"\d{1,5}", raw):
                    continue
                top = float(word.get("top", 0) or 0)
                x0 = float(word.get("x0", 0) or 0)
                x1 = float(word.get("x1", 0) or 0)
                vertical_gap = top - label_bottom
                if not (0 <= vertical_gap <= 45):
                    continue
                if x1 < label_x0 - 35 or x0 > label_x1 + 70:
                    continue
                candidates.append((vertical_gap, abs(x0 - label_x0), int(raw)))

            if candidates:
                candidates.sort()
                value = candidates[0][2]
                return value if value > 0 else None

    return None



def extract_spacing_from_page(
    text: str,
    words: Optional[list[dict[str, object]]] = None,
) -> Optional[float]:
    """
    Extrait la distance entre supports depuis une première page de support SUFIX.

    Variantes rencontrées et acceptées notamment :
    - « Espace entre les supports 1,5m »
    - « Espace entre les supports » puis « 1,5m » à la ligne suivante
    - « Espacement entre supports : 1,5 m »
    - « Espacement entre le support ... » / variantes singulier-pluriel

    La recherche reste volontairement locale au libellé pour éviter de capter
    d'autres longueurs de la nomenclature.
    """
    normalized = (text or "").replace("\u00a0", " ")
    # Normaliser les espaces horizontaux sans supprimer les retours de ligne.
    normalized = "\n".join(
        re.sub(r"[ \t]+", " ", line).strip()
        for line in normalized.splitlines()
    )

    label = r"(?:espace|espacement)\s+entre\s+(?:(?:les?|des?)\s+)?supports?"
    number = r"([0-9]+(?:[,.][0-9]+)?)"

    patterns = [
        # Même ligne : « Espace entre les supports 1,5m »
        rf"(?im)^\s*{label}\s*[:\-]?\s*{number}\s*m(?:ètres?)?\s*$",
        # Valeur après le libellé avec quelques caractères de mise en page.
        rf"(?is){label}\s*[:\-]?\s*.{{0,12}}?{number}\s*m\b",
        # Ligne suivante.
        rf"(?is){label}\s*[:\-]?\s*\n\s*{number}\s*m?\b",
    ]

    for pattern in patterns:
        match = re.search(pattern, normalized)
        if not match:
            continue
        try:
            value = float(match.group(1).replace(",", "."))
        except (TypeError, ValueError):
            continue
        # Une distance de supportage réaliste est strictement positive.
        # La limite large à 100 m protège surtout contre les faux positifs.
        if 0 < value < 100:
            return value

    # Recherche ligne par ligne, utile lorsque le moteur PDF sépare le libellé.
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    accepted_labels = {
        "espace entre support",
        "espace entre supports",
        "espace entre le support",
        "espace entre les supports",
        "espacement entre support",
        "espacement entre supports",
        "espacement entre le support",
        "espacement entre les supports",
    }

    for index, line in enumerate(lines):
        norm_line = normalize_text(line).strip(" :-")
        if norm_line not in accepted_labels:
            continue

        for candidate in lines[index + 1 : index + 3]:
            match = re.fullmatch(
                r"\s*([0-9]+(?:[,.][0-9]+)?)\s*m(?:ètres?)?\s*",
                candidate,
                re.I,
            )
            if match:
                value = float(match.group(1).replace(",", "."))
                if 0 < value < 100:
                    return value

            # Ne pas continuer jusqu'à la nomenclature si un nouveau champ est rencontré.
            if re.search(r"[A-Za-zÀ-ÿ]", candidate) and not re.fullmatch(
                r"\s*[0-9]+(?:[,.][0-9]+)?\s*m(?:ètres?)?\s*",
                candidate,
                re.I,
            ):
                break

    # Secours spatial. Les mots sont au format pdfplumber :
    # text, x0, x1, top, bottom.
    if words:
        ordered = sorted(
            words,
            key=lambda word: (
                float(word.get("top", 0) or 0),
                float(word.get("x0", 0) or 0),
            ),
        )

        label_tokens = [
            ("espace", "entre", "les", "supports"),
            ("espacement", "entre", "les", "supports"),
            ("espace", "entre", "supports"),
            ("espacement", "entre", "supports"),
        ]

        normalized_words = [
            normalize_text(str(word.get("text", ""))) for word in ordered
        ]

        for tokens in label_tokens:
            token_count = len(tokens)
            for idx in range(len(ordered) - token_count + 1):
                if tuple(normalized_words[idx : idx + token_count]) != tokens:
                    continue

                label_words = ordered[idx : idx + token_count]
                tops = [float(word.get("top", 0) or 0) for word in label_words]
                # Le libellé doit être sur une même ligne.
                if max(tops) - min(tops) > 5:
                    continue

                label_bottom = max(
                    float(word.get("bottom", 0) or 0) for word in label_words
                )
                label_x0 = min(float(word.get("x0", 0) or 0) for word in label_words)
                label_x1 = max(float(word.get("x1", 0) or 0) for word in label_words)

                candidates: list[tuple[float, float, float]] = []
                for word in ordered:
                    raw = str(word.get("text", "")).strip()
                    match = re.fullmatch(
                        r"([0-9]+(?:[,.][0-9]+)?)\s*m?",
                        raw,
                        re.I,
                    )
                    if not match:
                        continue

                    value = float(match.group(1).replace(",", "."))
                    if not (0 < value < 100):
                        continue

                    top = float(word.get("top", 0) or 0)
                    x0 = float(word.get("x0", 0) or 0)
                    x1 = float(word.get("x1", 0) or 0)

                    # Valeur sur la même ligne ou juste en dessous du libellé.
                    vertical_gap = top - label_bottom
                    same_line = abs(top - min(tops)) <= 5
                    just_below = -3 <= vertical_gap <= 45
                    if not (same_line or just_below):
                        continue

                    # Elle doit également rester dans une zone horizontale proche.
                    if x1 < label_x0 - 40 or x0 > label_x1 + 110:
                        continue

                    candidates.append(
                        (
                            0 if same_line else max(vertical_gap, 0),
                            abs(x0 - label_x0),
                            value,
                        )
                    )

                if candidates:
                    candidates.sort()
                    return candidates[0][2]

    return None


def _extract_support_information_pdfplumber(pdf_path: Path, rows: list[dict[str, object]]) -> list[SupportInfo]:
    expected_supports = ordered_unique_supports(rows)
    levels = {str(r.get("Support") or ""): str(r.get("Niveau") or "") for r in rows}
    infos: list[SupportInfo] = []
    cache: dict[int, str] = {}
    with pdfplumber.open(pdf_path) as pdf:
        pages = _extract_support_pages_from_pdf(pdf, expected_supports, cache)
        for support in expected_supports:
            page_number = pages.get(support)
            if page_number is None or not 1 <= page_number <= len(pdf.pages):
                infos.append(SupportInfo(support, None, 1, levels.get(support, ""), "PDF - support non localisé", None, "", True))
                continue
            page = pdf.pages[page_number-1]
            text = cache.get(page_number) or page.extract_text() or ""
            cache[page_number] = text
            count = extract_count_from_page(text)
            if count is None and "nombre de supports" in normalize_text(text):
                try: count = extract_count_from_page(text, page.extract_words() or [])
                except Exception: pass
            spacing = extract_spacing_from_page(text)
            normalized_page = normalize_text(text)
            if spacing is None and (
                "espace entre" in normalized_page
                or "espacement entre" in normalized_page
            ):
                try:
                    spacing = extract_spacing_from_page(
                        text,
                        page.extract_words() or [],
                    )
                except Exception:
                    pass
            chassis = count is None and spacing is None
            infos.append(SupportInfo(
                support, page_number, count or 1, levels.get(support, ""),
                "PDF - lecture stricte" if count else ("Châssis" if chassis else "PDF - valeur proposée"),
                spacing, "PDF - espacement" if spacing is not None else ("Châssis" if chassis else "Non détecté"), chassis
            ))
    return infos


def extract_support_information(pdf_path: Path, rows: list[dict[str, object]]) -> list[SupportInfo]:
    if fitz is None:
        return _extract_support_information_pdfplumber(pdf_path, rows)
    expected_supports = ordered_unique_supports(rows)
    levels = {str(r.get("Support") or ""): str(r.get("Niveau") or "") for r in rows}
    try:
        infos: list[SupportInfo] = []
        cache: dict[int, str] = {}
        with fitz.open(pdf_path) as document:
            pages = _extract_support_pages_fitz(document, expected_supports, cache)
            for support in expected_supports:
                page_number = pages.get(support)
                if page_number is None or not 1 <= page_number <= document.page_count:
                    infos.append(SupportInfo(support, None, 1, levels.get(support, ""), "PDF - support non localisé", None, "", True))
                    continue
                text = _fitz_page_text(document, page_number, cache)
                count = extract_count_from_page(text)
                if count is None and "nombre de supports" in normalize_text(text):
                    count = extract_count_from_page(text, _fitz_words_as_pdfplumber(document, page_number))
                spacing = extract_spacing_from_page(text)
                normalized_page = normalize_text(text)
                if spacing is None and (
                    "espace entre" in normalized_page
                    or "espacement entre" in normalized_page
                ):
                    spacing = extract_spacing_from_page(
                        text,
                        _fitz_words_as_pdfplumber(
                            document,
                            page_number,
                        ),
                    )
                chassis = count is None and spacing is None
                infos.append(SupportInfo(
                    support, page_number, count or 1, levels.get(support, ""),
                    "PDF - lecture stricte" if count else ("Châssis" if chassis else "PDF - valeur proposée"),
                    spacing, "PDF - espacement" if spacing is not None else ("Châssis" if chassis else "Non détecté"), chassis
                ))
        return infos
    except Exception:
        return _extract_support_information_pdfplumber(pdf_path, rows)


def is_length_item(row: dict[str, object]) -> bool:
    return all(parse_float(row.get(col)) is not None for col in LENGTH_COLUMNS)


def add_quantities(
    rows: list[dict[str, object]],
    support_counts: dict[str, int],
    support_spacings: Optional[dict[str, Optional[float]]] = None,
) -> list[dict[str, object]]:
    output: list[dict[str, object]] = []
    spacings = support_spacings or {}
    for original in rows:
        row = dict(original)
        support = str(row["Support"])
        count = int(support_counts.get(support, 1))
        if count <= 0:
            raise ValueError(f"Le nombre de supports de {support} doit être supérieur à 0.")
        quantity = parse_float(row["Quantité"]) or 0.0
        row["Nombre de supports"] = count
        row["Espacement entre supports (m)"] = spacings.get(support)
        row["Quantité pour 1 support"] = quantity / count
        output.append(row)
    return output


def grouping_key(row: dict[str, object]) -> tuple:
    code = normalize_text(row.get("Code article"))
    if is_length_item(row):
        return (
            code,
            round(parse_float(row.get("Longueur utile")) or 0.0, 6),
            round(parse_float(row.get("Longueur proposée")) or 0.0, 6),
            round(parse_float(row.get("Chute calculée")) or 0.0, 6),
        )
    return (code, None, None, None)


def aggregate_rows(
    rows: list[dict[str, object]],
    quantity_column: str,
    keep_columns: list[str],
) -> list[dict[str, object]]:
    groups: dict[tuple, dict[str, object]] = {}
    order: list[tuple] = []
    for row in rows:
        key = grouping_key(row)
        if key not in groups:
            groups[key] = {col: row.get(col) for col in keep_columns}
            groups[key][quantity_column] = 0.0
            order.append(key)
        groups[key][quantity_column] = (
            parse_float(groups[key].get(quantity_column)) or 0.0
        ) + (parse_float(row.get(quantity_column)) or 0.0)

    result = [groups[key] for key in order]
    result.sort(
        key=lambda row: (
            0 if is_length_item(row) else 1,
            normalize_text(row.get("Code article")),
            -(parse_float(row.get("Longueur utile")) or 0.0),
        )
    )
    return result


def build_by_support(rows: list[dict[str, object]]) -> list[tuple[str, list[dict[str, object]]]]:
    sections: list[tuple[str, list[dict[str, object]]]] = []
    supports = sorted({str(row["Support"]) for row in rows}, key=natural_number)
    kept = [
        "Code article",
        "Libellé",
        "Quantité pour 1 support",
        "Longueur utile",
        "Longueur proposée",
        "Chute calculée",
    ]
    for support in supports:
        support_rows = [row for row in rows if str(row["Support"]) == support]
        if not support_rows:
            continue
        level = support_rows[0].get("Niveau") or ""
        count = int(parse_float(support_rows[0].get("Nombre de supports")) or 1)
        spacing = parse_float(support_rows[0].get("Espacement entre supports (m)"))
        spacing_text = f" - Espacement : {spacing:g} m" if spacing is not None else " - Espacement : Châssis / N.A."
        title = f"{support} - {level} - Nombre de supports : {count}{spacing_text}"
        aggregated = aggregate_rows(
            support_rows,
            quantity_column="Quantité pour 1 support",
            keep_columns=kept,
        )
        sections.append((title, aggregated))
    return sections


def build_total_for_optimization(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    kept = [
        "Code article",
        "Libellé",
        "Quantité",
        "Longueur utile",
        "Longueur proposée",
        "Chute calculée",
    ]
    return aggregate_rows(rows, quantity_column="Quantité", keep_columns=kept)


@lru_cache(maxsize=16384)
def _normalize_code_cached(text: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", text.upper())


def normalize_code(code: object) -> str:
    return _normalize_code_cached(str(code or ""))


def _article_version_tuple(header: object) -> Optional[tuple[int, ...]]:
    """Détecte une colonne de type Code article V1.0.9 / Code_Article_V1.0.10."""
    normalized = normalize_text(header)
    match = re.search(
        r"code[_ ]*article[_ ]*v\s*(\d+(?:\.\d+)+)",
        normalized,
    )
    if not match:
        return None
    try:
        return tuple(int(part) for part in match.group(1).split("."))
    except ValueError:
        return None


def _is_old_article_code_header(header: object) -> bool:
    normalized = normalize_text(header)
    compact = re.sub(r"[ _-]+", "", normalized)
    return "oldcodearticle" in compact or "anciencodearticle" in compact


def _is_legacy_current_code_header(header: object) -> bool:
    normalized = normalize_text(header)
    compact = re.sub(r"[ _-]+", "", normalized)
    return compact == "codearticle"


def load_article_database(database_path: Path) -> dict[str, dict[str, object]]:
    """
    Charge la base article en tenant compte de tout l'historique des codes.

    Le fichier peut évoluer avec des colonnes successives :
    OLD_Code_Article -> Code article V1.0.9 -> Code article V1.0.10 -> ...

    Chaque ancien code est résolu étape par étape jusqu'à la dernière version.
    Le dictionnaire retourné permet donc de rechercher aussi bien un ancien code
    qu'un code intermédiaire ou le code courant.
    """
    if not database_path.exists():
        raise FileNotFoundError(f"Base article introuvable : {database_path}")

    wb = load_workbook(database_path, data_only=True, read_only=True)
    sheet_name = "Base article" if "Base article" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}

    raw_headers = ["" if value is None else str(value).strip() for value in rows[0]]
    normalized_headers = [normalize_text(value) for value in raw_headers]

    version_columns: list[tuple[tuple[int, ...], int, str]] = []
    old_columns: list[int] = []
    legacy_columns: list[int] = []
    for index, header in enumerate(raw_headers):
        version = _article_version_tuple(header)
        if version is not None:
            version_columns.append((version, index, header))
        elif _is_old_article_code_header(header):
            old_columns.append(index)
        elif _is_legacy_current_code_header(header):
            legacy_columns.append(index)

    version_columns.sort(key=lambda item: item[0])
    if not version_columns and not legacy_columns:
        raise ValueError(
            "Aucune colonne de code article n'a été détectée dans la base. "
            "Utilisez par exemple « Code article V1.0.9 »."
        )

    current_code_index = version_columns[-1][1] if version_columns else legacy_columns[-1]

    def index_of(*names: str) -> Optional[int]:
        targets = {normalize_text(name) for name in names}
        for idx, header in enumerate(normalized_headers):
            if header in targets:
                return idx
        return None

    label_index = index_of("Libellé Produit", "Libelle Produit", "Libellé", "Libelle")
    category_index = index_of("Category", "Catégorie", "Categorie")
    packaging_index = index_of("Packaging")
    unit_index = index_of("Unité de vente", "Unite de vente")

    missing = []
    if label_index is None: missing.append("Libellé Produit")
    if category_index is None: missing.append("Category")
    if packaging_index is None: missing.append("Packaging")
    if unit_index is None: missing.append("Unité de vente")
    if missing:
        raise ValueError(
            "Colonnes manquantes dans la base article : " + ", ".join(missing)
        )

    # Transitions successives entre générations de codes.
    transitions: dict[str, str] = {}
    transition_conflicts: dict[str, set[str]] = defaultdict(set)
    metadata_by_current: dict[str, dict[str, object]] = {}
    observed_aliases: set[str] = set()
    composite_metadata: dict[str, dict[str, object]] = {}

    ordered_code_columns = old_columns + [column[1] for column in version_columns]
    if not version_columns:
        ordered_code_columns += legacy_columns

    for values in rows[1:]:
        if not values:
            continue

        stage_values: list[str] = []
        for column_index in ordered_code_columns:
            value = values[column_index] if column_index < len(values) else None
            key = normalize_code(value)
            if key and (not stage_values or key != stage_values[-1]):
                stage_values.append(key)
                observed_aliases.add(key)

        current_value = values[current_code_index] if current_code_index < len(values) else None
        current_key = normalize_code(current_value)
        if not current_key:
            # Si la dernière version est vide sur une ligne, reprendre la version
            # non vide la plus récente de cette ligne.
            current_key = stage_values[-1] if stage_values else ""
            current_value = current_key
        if not current_key:
            continue
        observed_aliases.add(current_key)

        for older, newer in zip(stage_values, stage_values[1:]):
            if older != newer:
                transition_conflicts[older].add(newer)

        label = values[label_index] if label_index is not None and label_index < len(values) else None
        category = values[category_index] if category_index is not None and category_index < len(values) else None
        packaging = parse_float(values[packaging_index]) if packaging_index is not None and packaging_index < len(values) else None
        unit = values[unit_index] if unit_index is not None and unit_index < len(values) else None

        info = {
            "Code article": str(current_value).strip(),
            "Libellé": "" if label is None else str(label).strip(),
            "Category": "Non classé" if category in (None, "") else str(category).strip(),
            "Packaging": packaging if packaging and packaging > 0 else 1.0,
            "Unité de vente": "UN" if unit in (None, "") else str(unit).strip(),
        }
        metadata_by_current[current_key] = info

        label_key = normalize_text(label)
        for alias in stage_values or [current_key]:
            if label_key:
                composite_metadata[f"{alias}||{label_key}"] = info

    # Valider les transitions. Un même ancien code ne doit pas pointer vers deux
    # références différentes, sauf si les branches convergent ensuite vers le
    # même code final.
    for old_key, targets in transition_conflicts.items():
        if len(targets) == 1:
            transitions[old_key] = next(iter(targets))
        else:
            # On conserve les conflits pour la résolution par code+libellé et on
            # évite un remplacement silencieux potentiellement faux.
            pass

    def resolve(code_key: str) -> str:
        current = code_key
        seen: set[str] = set()
        while current in transitions and current not in seen:
            seen.add(current)
            current = transitions[current]
        return current

    database: dict[str, dict[str, object]] = {}
    for alias in observed_aliases:
        final_key = resolve(alias)
        info = metadata_by_current.get(final_key)
        if info is None:
            # Un code peut être la valeur courante d'une ligne dupliquée.
            info = metadata_by_current.get(alias)
        if info is not None:
            database[alias] = info

    # Les clés composites servent à désambiguïser d'éventuels anciens codes en
    # double (par exemple une valeur générique) grâce au libellé SUFIX.
    database.update(composite_metadata)
    wb.close()
    return database



def article_database_version(database_path: Path) -> str:
    try:
        wb = load_workbook(database_path, data_only=True, read_only=True)
        ws = wb["Base article"] if "Base article" in wb.sheetnames else wb[wb.sheetnames[0]]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        versions = [(v, str(h)) for h in headers if (v := _article_version_tuple(h)) is not None]
        return max(versions)[1].replace("Code article", "").strip() if versions else "Non versionnée"
    except Exception:
        return "Inconnue"


def is_cable_tray_category(category: object) -> bool:
    value = normalize_text(category)
    return value.startswith("chemin de cables") or value.startswith("couvercle de chemin de cable")


def is_cable_tray_code(code: object, article_database: dict[str, dict[str, object]]) -> bool:
    info = article_database.get(normalize_code(code), {})
    return is_cable_tray_category(info.get("Category"))


def build_cable_tray_rows(
    enriched_rows: list[dict[str, object]],
    article_database: dict[str, dict[str, object]],
) -> tuple[list[dict[str, object]], list[str]]:
    grouped: dict[str, dict[str, object]] = {}
    warnings: list[str] = []
    for row in enriched_rows:
        code = str(row.get("Code article") or "").strip()
        key = normalize_code(code)
        info = article_database.get(key, {})
        if not is_cable_tray_category(info.get("Category")):
            continue
        spacing = parse_float(row.get("Espacement entre supports (m)"))
        count = int(parse_float(row.get("Nombre de supports")) or 1)
        per_support = parse_float(row.get("Quantité pour 1 support")) or 0.0
        if spacing is None or spacing <= 0:
            warnings.append(f"{row.get('Support')} : espacement manquant pour {code}")
            meters = 0.0
        else:
            meters = spacing * (count + 1) * per_support
        if key not in grouped:
            grouped[key] = {
                "Code article": code,
                "Libellé": info.get("Libellé") or row.get("Libellé") or "",
                "Category": info.get("Category") or "Chemin de câbles",
                "Métré nécessaire (m)": 0.0,
                "Longueur commerciale (m)": CABLE_TRAY_STOCK_LENGTH_M,
                "Nombre de barres": 0,
                "Métré à chiffrer (m)": 0.0,
                "Chute théorique (m)": 0.0,
                "Unité": info.get("Unité de vente") or "ML",
                "Packaging": parse_float(info.get("Packaging")) or CABLE_TRAY_STOCK_LENGTH_M,
            }
        grouped[key]["Métré nécessaire (m)"] += meters
    result=[]
    for key,row in grouped.items():
        meters=float(row["Métré nécessaire (m)"])
        stock=float(row["Packaging"] or CABLE_TRAY_STOCK_LENGTH_M)
        if stock <= 0: stock=CABLE_TRAY_STOCK_LENGTH_M
        bars=math.ceil(meters/stock) if meters>0 else 0
        row["Longueur commerciale (m)"]=stock
        row["Nombre de barres"]=bars
        row["Métré à chiffrer (m)"]=bars*stock
        row["Chute théorique (m)"]=max(0.0,bars*stock-meters)
        result.append(row)
    result.sort(key=lambda r:(normalize_text(r.get("Category")),normalize_text(r.get("Code article"))))
    return result,warnings


def cable_tray_offer_rows(cable_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    return [{
        "Type":"Chemin de câbles",
        "Code article":r.get("Code article"),
        "Profil":"",
        "Libellé":r.get("Libellé"),
        "Category":r.get("Category"),
        "Quantité à chiffrer":r.get("Métré à chiffrer (m)"),
        "Unité":r.get("Unité") or "ML",
        "Packaging":r.get("Packaging") or CABLE_TRAY_STOCK_LENGTH_M,
        "Nombre de code à chiffrer":r.get("Nombre de barres") or 0,
    } for r in cable_rows]

def apply_article_database(
    rows: list[dict[str, object]],
    article_database: dict[str, dict[str, object]],
) -> tuple[list[dict[str, object]], dict[str, int]]:
    """
    Remplace codes et libellés immédiatement après lecture du CSV, avant toute
    rationalisation ou optimisation.
    """
    output: list[dict[str, object]] = []
    stats = {"codes_remplaces": 0, "libelles_remplaces": 0, "non_trouves": 0}

    for original in rows:
        row = dict(original)
        source_code = str(row.get("Code article") or "").strip()
        source_label = str(row.get("Libellé") or "").strip()
        code_key = normalize_code(source_code)
        composite_key = f"{code_key}||{normalize_text(source_label)}"
        info = article_database.get(composite_key) or article_database.get(code_key)

        if info is None:
            stats["non_trouves"] += 1
            output.append(row)
            continue

        new_code = str(info.get("Code article") or source_code).strip()
        new_label = str(info.get("Libellé") or source_label).strip()
        if normalize_code(new_code) != normalize_code(source_code):
            stats["codes_remplaces"] += 1
        if new_label and normalize_text(new_label) != normalize_text(source_label):
            stats["libelles_remplaces"] += 1

        row["Code article"] = new_code
        if new_label:
            row["Libellé"] = new_label
        output.append(row)

    return output, stats


def augment_profile_mapping_with_article_history(
    code_to_profile: dict[str, str],
    article_database: dict[str, dict[str, object]],
) -> dict[str, str]:
    """Préserve la reconnaissance des profils après mise à jour des codes article."""
    result = dict(code_to_profile)
    changed = True
    while changed:
        changed = False
        for alias_key, info in article_database.items():
            if "||" in alias_key:
                continue
            current_key = normalize_code(info.get("Code article"))
            if alias_key in result and current_key and current_key not in result:
                result[current_key] = result[alias_key]
                changed = True
            elif current_key in result and alias_key not in result:
                result[alias_key] = result[current_key]
                changed = True
    return result


def normalize_stock_articles_with_article_database(
    stock_articles: dict[tuple[str, float], dict[str, str]],
    article_database: dict[str, dict[str, object]],
) -> dict[tuple[str, float], dict[str, str]]:
    result: dict[tuple[str, float], dict[str, str]] = {}
    for key, article in stock_articles.items():
        code_value = article.get("Code article")
        info = article_database.get(normalize_code(code_value))
        if info is None:
            result[key] = dict(article)
        else:
            result[key] = {
                "Code article": str(info.get("Code article") or code_value or ""),
                "Libellé": str(info.get("Libellé") or article.get("Libellé") or ""),
            }
    return result


@lru_cache(maxsize=4)
def _load_profile_sheet_values(database_path_text: str) -> dict[str, tuple[tuple[object, ...], ...]]:
    database_path = Path(database_path_text)
    if not database_path.exists():
        raise FileNotFoundError(f"Base de données introuvable : {database_path}")
    wb = load_workbook(database_path, data_only=True, read_only=True)
    try:
        return {
            name: tuple(tuple(row) for row in wb[name].iter_rows(values_only=True))
            for name in wb.sheetnames
        }
    finally:
        wb.close()


def _profile_sheet_values(database_path: Path, name: str) -> tuple[tuple[object, ...], ...]:
    sheets = _load_profile_sheet_values(str(database_path.resolve()))
    if name not in sheets:
        raise ValueError(f"Feuille absente de la base : {name}")
    return sheets[name]


def load_magnelis_replacement_mapping(
    database_path: Path,
) -> dict[str, str]:
    """Associe chaque profil Magnelis à son profil EZ de géométrie équivalente."""
    try:
        values = list(_profile_sheet_values(database_path, "Profiles de rails"))
    except ValueError:
        return {}
    if not values:
        return {}

    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    indexes = {header: pos for pos, header in enumerate(headers)}
    for required in ("Key", "Label", "Material"):
        if required not in indexes:
            return {}

    ez_by_geometry: dict[str, str] = {}
    magnelis_rows: list[tuple[str, str]] = []

    def geometry_key(label: object) -> str:
        text = normalize_text(label)
        text = re.sub(r"\bmagnelis\b", "", text)
        text = re.sub(r"\belectrozingue\b", "", text)
        return re.sub(r"\s+", "", text)

    for row in values[1:]:
        key_value = parse_float(row[indexes["Key"]])
        if key_value is None:
            continue
        profile_id = f"RAIL:{int(key_value)}"
        label = str(row[indexes["Label"]] or "")
        material = normalize_text(row[indexes["Material"]])
        geometry = geometry_key(label)
        if "magnelis" in material or "magnelis" in normalize_text(label):
            magnelis_rows.append((profile_id, geometry))
        elif "electro" in material or "zing" in material:
            ez_by_geometry[geometry] = profile_id

    mapping: dict[str, str] = {}
    for source_profile, geometry in magnelis_rows:
        target = ez_by_geometry.get(geometry)
        if target:
            mapping[source_profile] = target
    return mapping


def apply_magnelis_replacement(
    rows: list[dict[str, object]],
    options: dict[str, ProfileOption],
    code_to_profile: dict[str, str],
    stock_articles: dict[tuple[str, float], dict[str, str]],
    replacement_mapping: dict[str, str],
) -> tuple[list[dict[str, object]], int]:
    """
    Remplace les rails Magnelis par le profil EZ équivalent.

    La longueur commerciale d'origine est conservée si elle existe en EZ.
    Sinon, la longueur disponible la plus proche capable de contenir la découpe
    est utilisée. Longueur utile et quantité restent inchangées.
    """
    result: list[dict[str, object]] = []
    replacement_count = 0

    for source_row in rows:
        row = dict(source_row)
        source_profile = code_to_profile.get(normalize_code(row.get("Code article")))
        target_profile = replacement_mapping.get(source_profile or "")
        if not target_profile:
            result.append(row)
            continue

        useful = parse_float(row.get("Longueur utile"))
        proposed = parse_float(row.get("Longueur proposée"))
        target_option = options.get(target_profile)
        if target_option is None or not target_option.stock_lengths_m:
            result.append(row)
            continue

        available = sorted(float(value) for value in target_option.stock_lengths_m)
        compatible = [length for length in available if useful is None or length + 1e-9 >= useful]
        if not compatible:
            result.append(row)
            continue

        if proposed is not None and any(math.isclose(length, proposed, abs_tol=1e-9) for length in compatible):
            chosen = next(length for length in compatible if math.isclose(length, proposed, abs_tol=1e-9))
        elif proposed is not None:
            chosen = min(compatible, key=lambda length: (abs(length - proposed), length))
        else:
            chosen = min(compatible)

        article = stock_articles.get((target_profile, round(chosen, 6)))
        if not article:
            result.append(row)
            continue

        row["Code article"] = str(article.get("Code article") or row.get("Code article") or "")
        row["Libellé"] = str(article.get("Libellé") or row.get("Libellé") or "")
        row["Longueur proposée"] = chosen
        if useful is not None:
            row["Chute calculée"] = max(0.0, chosen - useful)
        replacement_count += 1
        result.append(row)

    return result, replacement_count


def build_initial_offer_rows(
    total_rows: list[dict[str, object]],
    article_database: dict[str, dict[str, object]],
    code_to_profile: dict[str, str],
    options: dict[str, ProfileOption],
    cable_rows: Optional[list[dict[str, object]]] = None,
) -> list[dict[str, object]]:
    """Construit l'offre commerciale sans optimiser les longueurs de barres."""
    grouped: dict[str, dict[str, object]] = {}

    for row in total_rows:
        code_value = str(row.get("Code article") or "").strip()
        key = normalize_code(code_value)
        if not key:
            continue
        if cable_rows and is_cable_tray_code(code_value, article_database):
            continue
        if key not in grouped:
            grouped[key] = {
                "Code article": code_value,
                "Libellé": str(row.get("Libellé") or ""),
                "Quantité": 0.0,
                "Est longueur": False,
            }
        grouped[key]["Quantité"] = (parse_float(grouped[key]["Quantité"]) or 0.0) + (parse_float(row.get("Quantité")) or 0.0)
        if is_length_item(row):
            grouped[key]["Est longueur"] = True

    output: list[dict[str, object]] = []
    for key, row in grouped.items():
        info = article_database.get(key, {
            "Category": "Non classé",
            "Packaging": 1.0,
            "Unité de vente": "UN",
            "Libellé": row.get("Libellé") or "",
        })
        quantity = parse_float(row.get("Quantité")) or 0.0
        packaging = parse_float(info.get("Packaging")) or 1.0
        sale_unit = str(info.get("Unité de vente") or "UN").strip()
        profile_id = code_to_profile.get(key, "")
        profile_label = options[profile_id].profile_label if profile_id in options else ""

        if bool(row.get("Est longueur")) and normalize_text(sale_unit) in {
            "ml", "m", "metre lineaire", "metres lineaires"
        }:
            quantity_to_price = quantity * packaging
            item_type = "Barre"
        else:
            quantity_to_price = quantity
            item_type = "Pièce"

        output.append({
            "Type": item_type,
            "Code article": row.get("Code article"),
            "Profil": profile_label,
            "Libellé": row.get("Libellé") or info.get("Libellé") or "",
            "Category": info.get("Category") or "Non classé",
            "Quantité à chiffrer": quantity_to_price,
            "Unité": sale_unit,
            "Packaging": packaging,
            "Nombre de code à chiffrer": math.ceil(quantity_to_price / packaging),
        })

    if cable_rows:
        output.extend(cable_tray_offer_rows(cable_rows))
    output.sort(key=lambda row: (normalize_text(row.get("Category")), normalize_text(row.get("Code article"))))
    return output


def load_stock_article_mapping(
    database_path: Path,
) -> dict[tuple[str, float], dict[str, str]]:
    """
    Associe un profil de rail/tige et une longueur commerciale à son code article.
    """
    mapping: dict[tuple[str, float], dict[str, str]] = {}

    definitions = [
        ("Codes rails", "RailProfileKey", "RAIL"),
        ("Codes tiges", "FastenerProfileKey", "TIGE"),
    ]
    for sheet_name, key_column, prefix in definitions:
        try:
            rows = list(_profile_sheet_values(database_path, sheet_name))
        except ValueError:
            continue
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        indexes = {header: pos for pos, header in enumerate(headers)}
        for required in ("Code article", "Label", key_column, "Length"):
            if required not in indexes:
                raise ValueError(f"Colonne {required} absente de {sheet_name}.")
        for values in rows[1:]:
            key_value = parse_float(values[indexes[key_column]])
            length = parse_float(values[indexes["Length"]])
            code_value = values[indexes["Code article"]]
            if key_value is None or length is None or not code_value:
                continue
            profile_id = f"{prefix}:{int(key_value)}"
            mapping[(profile_id, round(length, 6))] = {
                "Code article": str(code_value).strip(),
                "Libellé": str(values[indexes["Label"]] or "").strip(),
            }
    return mapping


def enrich_summary_rows(
    cut_summary: list[dict[str, object]],
    piece_rows: list[dict[str, object]],
    article_database: dict[str, dict[str, object]],
) -> list[dict[str, object]]:
    summary_rows: list[dict[str, object]] = []

    def article_info(code_value: object) -> dict[str, object]:
        return article_database.get(
            normalize_code(code_value),
            {
                "Code article": "" if code_value is None else str(code_value),
                "Libellé": "",
                "Category": "Non classé",
                "Packaging": 1.0,
                "Unité de vente": "UN",
            },
        )

    for row in cut_summary:
        code_value = row.get("Code article")
        info = article_info(code_value)
        number_of_bars = parse_float(row.get("Nombre de barres")) or 0.0
        packaging = parse_float(info.get("Packaging")) or 1.0
        sale_unit = str(info.get("Unité de vente") or "UN").strip()

        # Pour les articles vendus au mètre linéaire, le besoin à chiffrer
        # est exprimé en ML : nombre de barres x longueur commerciale.
        # Le Packaging correspond alors à la longueur contenue dans un code.
        if normalize_text(sale_unit) in {"ml", "m", "metre lineaire", "metre lineaire"}:
            quantity_to_price = number_of_bars * packaging
        else:
            quantity_to_price = number_of_bars

        summary_rows.append(
            {
                "Type": "Barre",
                "Code article": code_value,
                "Profil": row.get("Profil"),
                "Libellé": info.get("Libellé") or row.get("Libellé"),
                "Category": info.get("Category"),
                "Quantité à chiffrer": quantity_to_price,
                "Unité": sale_unit,
                "Packaging": packaging,
                "Nombre de code à chiffrer": math.ceil(quantity_to_price / packaging),
            }
        )

    for row in piece_rows:
        code_value = row.get("Code article")
        info = article_info(code_value)
        quantity = parse_float(row.get("Quantité")) or 0.0
        packaging = parse_float(info.get("Packaging")) or 1.0
        sale_unit = str(info.get("Unité de vente") or "UN").strip()
        summary_rows.append(
            {
                "Type": "Pièce",
                "Code article": code_value,
                "Profil": "",
                "Libellé": row.get("Libellé") or info.get("Libellé"),
                "Category": info.get("Category"),
                "Quantité à chiffrer": quantity,
                "Unité": sale_unit,
                "Packaging": packaging,
                "Nombre de code à chiffrer": math.ceil(quantity / packaging),
            }
        )

    summary_rows.sort(
        key=lambda row: (
            normalize_text(row.get("Category")),
            normalize_text(row.get("Code article")),
        )
    )
    return summary_rows


ERP_CSV_HEADERS = [
    "Numéro de Ligne",
    "Regroupement",
    "Type de ligne",
    "Code article",
    "Nom Article",
    "Statut article",
    "Qté demandée",
    "Unité de la qté demandée",
    "Qté Vendue",
    "Unité de la qté vendue",
    "Classe tarifaire",
    "Prix public",
    "R1%",
    "R2%",
    "R3%",
    "RA%",
    "Remise de Pied de page",
    "Prix net manuel",
    "Prix net unitaire",
    "Cout de revient unitaire",
    "Cout total incluant les frais",
    "Prix net total",
    "Prix total frais externes inclus",
    "Marge",
    "Marge %",
    "Impression Photo",
    "Total ECO",
    "Stock disponible",
    "Délai Horizon",
    "Désignation longue",
    "Texte article",
    "Pays origine",
    "Code douanier",
    "Mulitiple de sortie",
    "Repérage",
    "Texte libre 1",
    "Texte libre 2",
    "Catalogue",
    "Type article",
    "Hors gabarit",
    "Poids brut",
    "Planification",
    "Classe ABC",
    "Texte description longue article",
    "Texte description courte article",
    "Description courte",
    "Commentaires",
]


def _erp_text(value: object) -> str:
    """Reproduit Trim(Replace(CStr(Cel.Text), ",", "."))."""
    if value is None:
        return ""
    if isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, int):
        text = str(value)
    elif isinstance(value, float):
        if math.isclose(value, round(value), abs_tol=1e-9):
            text = str(int(round(value)))
        else:
            text = format(value, ".12g")
    else:
        text = str(value)
    return text.strip().replace(",", ".").replace("\r", " ").replace("\n", " ")


def export_to_eq_csv(
    destination_path: Path,
    summary_rows: list[dict[str, object]],
) -> None:
    """
    Génère directement le CSV final accepté par EQ.

    - UTF-8 sans BOM ;
    - séparateur virgule ;
    - 47 colonnes dans le même ordre que la macro historique ;
    - virgules internes remplacées par des points ;
    - regroupement par Category avec ligne de commentaire avant chaque groupe.
    """
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in summary_rows:
        grouped[str(row.get("Category") or "Non classé")].append(row)

    output_lines = [",".join(ERP_CSV_HEADERS)]
    line_number = 1

    for category in sorted(grouped, key=normalize_text):
        comment_row = [""] * len(ERP_CSV_HEADERS)
        comment_row[0] = line_number
        comment_row[1] = category
        comment_row[2] = ERP_COMMENT_LINE_TYPE
        comment_row[46] = category
        output_lines.append(",".join(_erp_text(value) for value in comment_row))
        line_number += 1

        for item in sorted(
            grouped[category],
            key=lambda row: normalize_text(row.get("Code article")),
        ):
            data_row = [""] * len(ERP_CSV_HEADERS)
            data_row[0] = line_number
            data_row[1] = category
            data_row[2] = "STD"
            data_row[3] = item.get("Code article")
            data_row[4] = item.get("Libellé")

            sale_unit = str(item.get("Unité") or "UN").strip()
            normalized_sale_unit = normalize_text(sale_unit)
            if normalized_sale_unit in {
                "ml",
                "m",
                "metre lineaire",
                "metres lineaires",
            }:
                erp_quantity = parse_float(item.get("Quantité à chiffrer")) or 0
            else:
                erp_quantity = parse_float(item.get("Nombre de code à chiffrer")) or 0

            data_row[6] = display_number(erp_quantity)
            data_row[7] = sale_unit
            output_lines.append(",".join(_erp_text(value) for value in data_row))
            line_number += 1

    # Path.write_text(encoding="utf-8") produit bien un UTF-8 sans BOM.
    destination_path.write_text("\n".join(output_lines) + "\n", encoding="utf-8")


def load_profile_database(database_path: Path) -> tuple[dict[str, ProfileOption], dict[str, str]]:
    if not database_path.exists():
        raise FileNotFoundError(
            f"Base de données introuvable : {database_path}"
        )
    def sheet_rows(name: str) -> tuple[list[str], list[dict[str, object]]]:
        values = list(_profile_sheet_values(database_path, name))
        if not values:
            return [], []
        headers = [str(v).strip() if v is not None else "" for v in values[0]]
        records = [
            {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
            for row in values[1:]
            if any(cell is not None for cell in row)
        ]
        return headers, records

    _, rail_codes = sheet_rows("Codes rails")
    _, rail_profiles = sheet_rows("Profiles de rails")
    _, rod_codes = sheet_rows("Codes tiges")
    _, rod_profiles = sheet_rows("Profiles de tiges")
    _, console_codes = sheet_rows("Codes consoles")
    _, console_profiles = sheet_rows("Profiles de consoles")

    rail_profile_labels = {
        str(int(parse_float(r.get("Key")) or 0)): str(r.get("Label") or "")
        for r in rail_profiles
    }
    rod_profile_labels = {
        str(int(parse_float(r.get("Key")) or 0)): str(r.get("Label") or "")
        for r in rod_profiles
    }
    console_profile_labels = {
        str(int(parse_float(r.get("Key")) or 0)): str(r.get("Label") or "")
        for r in console_profiles
    }

    profile_lengths: dict[tuple[str, str], set[float]] = defaultdict(set)
    code_to_profile: dict[str, str] = {}
    profile_labels: dict[str, str] = {}

    for record in rail_codes:
        key = str(int(parse_float(record.get("RailProfileKey")) or 0))
        profile_id = f"RAIL:{key}"
        length = parse_float(record.get("Length"))
        if length:
            profile_lengths[("Rail", profile_id)].add(length)
        code = normalize_code(record.get("Code article"))
        if code:
            code_to_profile[code] = profile_id
        profile_labels[profile_id] = rail_profile_labels.get(key) or str(record.get("Label") or profile_id)

    for record in rod_codes:
        key = str(int(parse_float(record.get("FastenerProfileKey")) or 0))
        profile_id = f"TIGE:{key}"
        length = parse_float(record.get("Length"))
        if length:
            profile_lengths[("Tige", profile_id)].add(length)
        code = normalize_code(record.get("Code article"))
        if code:
            code_to_profile[code] = profile_id
            # SUFIX ajoute parfois le préfixe PFD devant le code de la base.
            code_to_profile["PFD" + code] = profile_id
        profile_labels[profile_id] = rod_profile_labels.get(key) or str(record.get("Label") or profile_id)

    # Les consoles sont des articles à longueur fixe : elles sont identifiées
    # grâce à la base, mais elles ne possèdent aucune longueur de barre à choisir
    # et ne passent jamais dans l'algorithme d'optimisation de découpe.
    for record in console_codes:
        key = str(int(parse_float(record.get("ConsoleProfileKey")) or 0))
        profile_id = f"CONSOLE:{key}"
        code = normalize_code(record.get("Code article"))
        if code:
            code_to_profile[code] = profile_id
        profile_labels[profile_id] = (
            console_profile_labels.get(key)
            or str(record.get("Label") or profile_id)
        )

    options: dict[str, ProfileOption] = {}
    for (category, profile_id), lengths in profile_lengths.items():
        options[profile_id] = ProfileOption(
            profile_key=profile_id,
            profile_label=profile_labels.get(profile_id, profile_id),
            category=category,
            stock_lengths_m=tuple(sorted(lengths)),
        )

    # Ajout des profils de consoles sans longueurs disponibles.
    for profile_id, label in profile_labels.items():
        if profile_id.startswith("CONSOLE:"):
            options[profile_id] = ProfileOption(
                profile_key=profile_id,
                profile_label=label,
                category="Console",
                stock_lengths_m=tuple(),
            )

    return options, code_to_profile


def infer_profile_from_label(
    code: object,
    label: object,
    options: dict[str, ProfileOption],
    code_to_profile: dict[str, str],
) -> Optional[str]:
    normalized_code = normalize_code(code)
    candidates = [
        normalized_code,
        normalized_code.removeprefix("PFD"),
    ]
    for candidate in candidates:
        if candidate in code_to_profile:
            return code_to_profile[candidate]

    text = normalize_text(label)
    # Secours pour tiges filetées.
    diameter_match = re.search(r"\bm\s*(6|8|10|12|16)\b", text)
    if not diameter_match:
        diameter_match = re.search(r"diametre\s*(6|8|10|12|16)", text)
    if not diameter_match and "tige filetee" in text:
        code_match = re.search(r"(?:n|m)(06|08|10|12|16)", normalize_code(code))
        if code_match:
            diameter_match_value = str(int(code_match.group(1)))
        else:
            diameter_match_value = None
    else:
        diameter_match_value = diameter_match.group(1) if diameter_match else None

    if diameter_match_value:
        target = f"m{diameter_match_value}"
        for profile_id, option in options.items():
            if option.category == "Tige" and target in normalize_text(option.profile_label).replace(" ", ""):
                return profile_id

    # Secours pour rails : dimensions + épaisseur.
    dims = re.search(r"41\s*[x/_]\s*(21|41|62|82)", text)
    thickness = re.search(r"(1[,.]5|1[,.]8|2[,.]5|2)\s*mm", text)
    if dims and thickness:
        dim2 = dims.group(1)
        ep = thickness.group(1).replace(",", ".")
        for profile_id, option in options.items():
            normalized_label = normalize_text(option.profile_label).replace(",", ".")
            if f"41x{dim2}" in normalized_label.replace(" ", "") and ep in normalized_label:
                return profile_id
    return None


def attach_profiles(
    length_rows: list[dict[str, object]],
    options: dict[str, ProfileOption],
    code_to_profile: dict[str, str],
) -> tuple[list[dict[str, object]], list[str]]:
    enriched: list[dict[str, object]] = []
    unmapped: list[str] = []
    for row in length_rows:
        new_row = dict(row)
        profile_id = infer_profile_from_label(
            row.get("Code article"), row.get("Libellé"), options, code_to_profile
        )
        if profile_id:
            option = options[profile_id]
            new_row["Profil"] = option.profile_label
            new_row["_profile_id"] = profile_id
            new_row["Catégorie profil"] = option.category
        else:
            new_row["Profil"] = "Profil non reconnu"
            new_row["_profile_id"] = ""
            new_row["Catégorie profil"] = ""
            unmapped.append(f"{row.get('Code article')} - {row.get('Libellé')}")
        enriched.append(new_row)
    return enriched, unmapped


def cut_length_bounds(
    length_rows: list[dict[str, object]],
) -> dict[str, tuple[float, float]]:
    """Retourne la plus petite et la plus grande découpe de chaque profil."""
    bounds: dict[str, tuple[float, float]] = {}
    for row in length_rows:
        profile_id = str(row.get("_profile_id") or "")
        length_m = parse_float(row.get("Longueur utile"))
        if not profile_id or length_m is None or length_m <= 0:
            continue
        if profile_id not in bounds:
            bounds[profile_id] = (length_m, length_m)
        else:
            minimum, maximum = bounds[profile_id]
            bounds[profile_id] = (min(minimum, length_m), max(maximum, length_m))
    return bounds



def expand_cuts(length_rows: list[dict[str, object]]) -> dict[str, list[int]]:
    cuts_by_profile: dict[str, list[int]] = defaultdict(list)
    for row in length_rows:
        profile_id = str(row.get("_profile_id") or "")
        if not profile_id:
            continue
        length_m = parse_float(row.get("Longueur utile"))
        quantity = parse_float(row.get("Quantité")) or 0.0
        if length_m is None or length_m <= 0:
            continue

        rounded_quantity = round(quantity)
        if abs(quantity - rounded_quantity) > 1e-6:
            raise ValueError(
                f"La quantité de l'article {row.get('Code article')} n'est pas entière "
                f"({quantity}). L'optimisation de découpes exige un nombre entier de pièces."
            )
        cut_mm = int(round(length_m * 1000))
        cuts_by_profile[profile_id].extend([cut_mm] * int(rounded_quantity))
    return cuts_by_profile


def _clone_bins(bins: list[CutBin]) -> list[CutBin]:
    return [CutBin(bin.stock_length_mm, list(bin.cuts_mm)) for bin in bins]


def _reserve_exact_stock_lengths(
    cuts_mm: list[int],
    allowed_stock_lengths_mm: list[int],
) -> tuple[list[CutBin], list[int]]:
    """
    Affecte d'abord une découpe à une barre commerciale exactement de même
    longueur. Cela correspond notamment aux rails prédécoupés de 1 m.
    """
    allowed = set(allowed_stock_lengths_mm)
    exact_bins: list[CutBin] = []
    remaining: list[int] = []
    for cut in cuts_mm:
        if cut in allowed:
            exact_bins.append(CutBin(cut, [cut]))
        else:
            remaining.append(cut)
    return exact_bins, remaining


def _plan_mixed_lengths(
    cuts_mm: list[int],
    allowed_stock_lengths_mm: list[int],
    policy: str,
) -> list[CutBin]:
    """
    Construit un plan de découpe avec plusieurs longueurs commerciales.

    Plusieurs politiques sont testées ensuite et la meilleure est retenue :
    - smallest : ouvre la plus petite barre compatible ;
    - largest : ouvre la plus grande barre compatible, utile pour mutualiser
      plusieurs découpes dans une même barre ;
    - best_pack : choisit la longueur dont le remplissage immédiat est le meilleur.
    """
    remaining = sorted(cuts_mm, reverse=True)
    bins: list[CutBin] = []

    while remaining:
        cut = remaining.pop(0)

        # Utiliser en priorité une chute déjà ouverte.
        eligible = [bin for bin in bins if bin.waste_mm >= cut]
        if eligible:
            target = min(eligible, key=lambda bin: (bin.waste_mm - cut, bin.stock_length_mm))
            target.cuts_mm.append(cut)
            continue

        candidates = [
            stock for stock in allowed_stock_lengths_mm if stock >= cut
        ]
        if not candidates:
            raise ValueError(
                f"Aucune longueur de barre sélectionnée ne peut recevoir "
                f"une découpe de {cut / 1000:g} m."
            )

        if policy == "smallest":
            stock_length = min(candidates)
        elif policy == "largest":
            stock_length = max(candidates)
        else:
            # Simuler le remplissage de chaque longueur avec les découpes encore
            # disponibles et retenir celle qui laisse le moins de chute relative.
            simulated = []
            for stock in candidates:
                capacity = stock - cut
                used_extra = 0
                packed_count = 1
                for candidate_cut in remaining:
                    if candidate_cut <= capacity:
                        capacity -= candidate_cut
                        used_extra += candidate_cut
                        packed_count += 1
                used = cut + used_extra
                waste = stock - used
                simulated.append(
                    (
                        waste / stock,
                        -packed_count,
                        waste,
                        stock,
                    )
                )
            stock_length = min(simulated)[3]

        bins.append(CutBin(stock_length, [cut]))

    return bins


def optimize_cuts_multi(
    cuts_mm: list[int],
    allowed_stock_lengths_mm: list[int],
    prioritize_exact: bool = True,
    optimization_mode: str = "equilibre",
) -> list[CutBin]:
    """
    Optimisation multi-longueurs déterministe.

    Objectif métier :
    1. si demandé, utiliser les longueurs commerciales exactement égales aux
       découpes (par exemple une pièce de 1 m sur un rail de 1 m) ;
    2. réutiliser les chutes ouvertes avant d'ouvrir une nouvelle barre ;
    3. tester plusieurs stratégies de longueurs commerciales ;
    4. retenir le plan qui minimise la longueur totale achetée puis, à égalité,
       le nombre de barres.
    """
    stock_lengths = sorted(set(int(value) for value in allowed_stock_lengths_mm))
    if not stock_lengths:
        raise ValueError("Aucune longueur de barre n'a été sélectionnée.")
    if any(cut > max(stock_lengths) for cut in cuts_mm):
        longest = max(cuts_mm)
        raise ValueError(
            f"Une découpe de {longest / 1000:g} m dépasse toutes les "
            "longueurs de barres sélectionnées."
        )

    exact_bins: list[CutBin] = []
    remaining = sorted(cuts_mm, reverse=True)
    if prioritize_exact:
        exact_bins, remaining = _reserve_exact_stock_lengths(
            remaining,
            stock_lengths,
        )

    candidate_plans: list[list[CutBin]] = []
    for policy in ("smallest", "largest", "best_pack"):
        planned = _plan_mixed_lengths(remaining, stock_lengths, policy)
        candidate_plans.append(_clone_bins(exact_bins) + planned)

    def plan_score(bins: list[CutBin]) -> tuple:
        total_stock = sum(bin.stock_length_mm for bin in bins)
        total_waste = sum(bin.waste_mm for bin in bins)
        patterns = len({(b.stock_length_mm, tuple(sorted(b.cuts_mm, reverse=True))) for b in bins})
        mode = normalize_text(optimization_mode)
        if "fabrication" in mode or "simple" in mode:
            return patterns, len(bins), total_stock, total_waste
        if "matiere" in mode or "econom" in mode:
            return total_stock, total_waste, len(bins), patterns
        return total_stock, len(bins), patterns, total_waste

    best = min(candidate_plans, key=plan_score)
    best.sort(key=lambda bin: (bin.stock_length_mm, bin.waste_mm, -bin.used_mm))
    return best


def make_cut_results(
    length_rows: list[dict[str, object]],
    selected_lengths_m: dict[str, list[float]],
    options: dict[str, ProfileOption],
    stock_articles: dict[tuple[str, float], dict[str, str]],
    prioritize_exact: bool = True,
    optimization_mode: str = "equilibre",
) -> tuple[list[dict[str, object]], list[dict[str, object]]] :
    cuts_by_profile = expand_cuts(length_rows)
    summary: list[dict[str, object]] = []
    detail: list[dict[str, object]] = []

    for profile_id, cuts in sorted(
        cuts_by_profile.items(),
        key=lambda item: normalize_text(options[item[0]].profile_label),
    ):
        option = options[profile_id]
        selected = sorted(set(float(value) for value in selected_lengths_m[profile_id]))
        bins = optimize_cuts_multi(
            cuts,
            [int(round(value * 1000)) for value in selected],
            prioritize_exact=prioritize_exact,
            optimization_mode=optimization_mode,
        )

        bins_by_stock: dict[int, list[CutBin]] = defaultdict(list)
        for cut_bin in bins:
            bins_by_stock[cut_bin.stock_length_mm].append(cut_bin)

        for stock_length_mm, stock_bins in sorted(bins_by_stock.items()):
            stock_length_m = stock_length_mm / 1000
            stock_article = stock_articles.get(
                (profile_id, round(stock_length_m, 6)),
                {"Code article": "", "Libellé": option.profile_label},
            )
            total_used = sum(bin.used_mm for bin in stock_bins)
            total_stock = sum(bin.stock_length_mm for bin in stock_bins)
            summary.append(
                {
                    "Code article": stock_article.get("Code article"),
                    "Profil": option.profile_label,
                    "Libellé": stock_article.get("Libellé"),
                    "Catégorie": option.category,
                    "Taille des barres (m)": stock_length_m,
                    "Nombre de barres": len(stock_bins),
                    "Longueur utile totale (m)": total_used / 1000,
                    "Chute totale (m)": (total_stock - total_used) / 1000,
                    "Taux d'utilisation": total_used / total_stock if total_stock else 0,
                }
            )

        for index, cut_bin in enumerate(bins, start=1):
            detail.append(
                {
                    "Profil": option.profile_label,
                    "Barre n°": index,
                    "Taille de barre (m)": cut_bin.stock_length_mm / 1000,
                    "Découpes (m)": " + ".join(
                        f"{cut / 1000:g}" for cut in cut_bin.cuts_mm
                    ),
                    "Longueur utilisée (m)": cut_bin.used_mm / 1000,
                    "Chute (m)": cut_bin.waste_mm / 1000,
                    "Taux d'utilisation": (
                        cut_bin.used_mm / cut_bin.stock_length_mm
                        if cut_bin.stock_length_mm
                        else 0
                    ),
                }
            )

    return summary, detail




def build_optimization_kpis(
    length_rows: list[dict[str, object]],
    cut_summary: list[dict[str, object]],
    cut_detail: list[dict[str, object]],
) -> dict[str, float]:
    initial_stock = sum((parse_float(r.get("Quantité")) or 0) * (parse_float(r.get("Longueur proposée")) or 0) for r in length_rows)
    initial_bars = sum(parse_float(r.get("Quantité")) or 0 for r in length_rows)
    optimized_stock = sum((parse_float(r.get("Nombre de barres")) or 0) * (parse_float(r.get("Taille des barres (m)")) or 0) for r in cut_summary)
    optimized_bars = sum(parse_float(r.get("Nombre de barres")) or 0 for r in cut_summary)
    used = sum(parse_float(r.get("Longueur utilisée (m)")) or 0 for r in cut_detail)
    waste = sum(parse_float(r.get("Chute (m)")) or 0 for r in cut_detail)
    patterns = len({(r.get("Profil"),r.get("Taille de barre (m)"),r.get("Découpes (m)")) for r in cut_detail})
    return {
        "Barres initiales": initial_bars,
        "Barres optimisées": optimized_bars,
        "Longueur initiale (m)": initial_stock,
        "Longueur optimisée (m)": optimized_stock,
        "Gain longueur (m)": initial_stock-optimized_stock,
        "Gain longueur (%)": ((initial_stock-optimized_stock)/initial_stock*100) if initial_stock else 0,
        "Longueur utilisée (m)": used,
        "Chute optimisée (m)": waste,
        "Rendement matière (%)": (used/optimized_stock*100) if optimized_stock else 100,
        "Plans de découpe distincts": patterns,
    }


def build_control_rows(
    enriched_rows: list[dict[str, object]], article_database: dict[str, dict[str, object]],
    length_rows: list[dict[str, object]], unmapped_profiles: list[str],
    cable_rows: list[dict[str, object]], cable_warnings: list[str], cut_detail: list[dict[str, object]],
    base_article_version: str, base_profiles_version: str,
) -> list[dict[str, object]]:
    supports={str(r.get("Support") or "") for r in enriched_rows}
    bad_counts={s for s in supports if not any((parse_float(r.get("Nombre de supports")) or 0)>0 for r in enriched_rows if str(r.get("Support") or "")==s)}
    unknown_codes=sorted({str(r.get("Code article") or "") for r in enriched_rows if normalize_code(r.get("Code article")) not in article_database})
    invalid_pack=sorted({str(r.get("Code article") or "") for r in enriched_rows if (info:=article_database.get(normalize_code(r.get("Code article")))) and (parse_float(info.get("Packaging")) or 0)<=0})
    rows=[]
    def add(name,status,detail): rows.append({"Contrôle":name,"Statut":status,"Détail":detail})
    add("Supports", "OK" if not bad_counts else "ERREUR", f"{len(supports)-len(bad_counts)}/{len(supports)} nombres valides")
    add("Codes articles", "OK" if not unknown_codes else "ATTENTION", "Tous reconnus" if not unknown_codes else ", ".join(unknown_codes[:12]))
    add("Profils rails/tiges", "OK" if not unmapped_profiles else "ATTENTION", "Tous reconnus" if not unmapped_profiles else ", ".join(unmapped_profiles[:12]))
    add("Packaging", "OK" if not invalid_pack else "ERREUR", "Tous valides" if not invalid_pack else ", ".join(invalid_pack[:12]))
    add("Chemins de câbles", "OK" if not cable_warnings else "ERREUR", f"{len(cable_rows)} référence(s) calculée(s)" if not cable_warnings else " ; ".join(cable_warnings[:8]))
    requested=sum(int(round(parse_float(r.get("Quantité")) or 0)) for r in length_rows)
    assigned=sum(len(str(r.get("Découpes (m)") or "").split(" + ")) if r.get("Découpes (m)") else 0 for r in cut_detail)
    add("Découpes affectées", "OK" if requested==assigned else "ERREUR", f"{assigned}/{requested} découpes")
    add("Version base article", "INFO", base_article_version)
    add("Version base profils", "INFO", base_profiles_version)
    return rows

def add_excel_table(
    ws,
    start_row: int,
    headers: list[str],
    rows: list[dict[str, object]],
    name: str,
) -> int:
    """
    Écrit un tableau visuel sans créer d'objet Excel « Table » (ListObject).

    Les anciennes versions créaient des fichiers /xl/tables/table*.xml qui
    pouvaient être réparés par Excel à l'ouverture. Le rendu reste identique,
    mais le classeur ne dépend plus de ces objets XML fragiles.
    """
    text_columns = {
        "Niveau",
        "Support",
        "Code article",
        "Libellé",
        "Profil",
        "Catégorie",
        "Catégorie profil",
        "Category",
        "Type",
        "Code article / Profil",
        "Unité",
        "Découpes (m)",
    }

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, col_index, header)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = TABLE_BORDER

    for row_offset, row in enumerate(rows, start=1):
        row_index = start_row + row_offset
        for col_index, header in enumerate(headers, start=1):
            value = row.get(header)
            if header in text_columns:
                value = "" if value is None else str(value)
            else:
                value = display_number(value)

            cell = ws.cell(row_index, col_index, value)
            cell.border = TABLE_BORDER
            cell.alignment = BODY_ALIGNMENT

            if header in text_columns:
                cell.number_format = "@"
                cell.data_type = "s"
            elif header == "Taux d'utilisation" and isinstance(value, (int, float)):
                cell.number_format = "0.0%"
            elif isinstance(value, float):
                cell.number_format = "0.###"

        # Léger effet de bandes alternées, sans objet Table Excel.
        if row_offset % 2 == 0:
            for col_index in range(1, len(headers) + 1):
                ws.cell(row_index, col_index).fill = ALT_ROW_FILL

    return start_row + max(len(rows), 1)


def format_sheet(ws, widths: Optional[dict[str, float]] = None) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    default_widths = {
        "A": 20, "B": 42, "C": 20, "D": 18, "E": 19, "F": 18,
        "G": 20, "H": 22, "I": 22, "J": 22,
    }
    for col, width in (widths or default_widths).items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical=cell.alignment.vertical or "top",
                horizontal=cell.alignment.horizontal,
                wrap_text=True,
            )


def configure_print_layout(ws) -> None:
    """Prépare une feuille visible pour une impression propre sur une ou plusieurs pages."""
    if ws.max_row < 1 or ws.max_column < 1:
        return
    ws.page_setup.orientation = "landscape" if ws.max_column >= 6 else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.page_margins.header = 0.20
    ws.page_margins.footer = 0.20
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.oddHeader.center.text = "SUFIX OptimiZer"
    ws.oddHeader.center.size = 9
    ws.oddFooter.center.text = "Page &[Page] / &[Pages]"
    ws.oddFooter.center.size = 8
    if ws.max_row > 1 and ws.title != SHEET_BY_SUPPORT:
        ws.print_title_rows = "1:1"


def write_global_sheet(ws, rows: list[dict[str, object]]) -> None:
    add_excel_table(ws, 1, OUTPUT_GLOBAL_COLUMNS, rows, "TableEtudeGlobale")
    format_sheet(ws)


def write_by_support_sheet(
    ws,
    sections: list[tuple[str, list[dict[str, object]]]],
) -> None:
    ws.sheet_view.showGridLines = False
    headers = [
        "Code article",
        "Libellé",
        "Quantité pour 1 support",
        "Longueur utile",
        "Longueur proposée",
        "Chute calculée",
    ]
    current_row = 1
    for index, (title, rows) in enumerate(sections, start=1):
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
        title_cell = ws.cell(current_row, 1, title)
        title_cell.fill = TITLE_FILL
        title_cell.font = TITLE_FONT
        title_cell.alignment = Alignment(vertical="center")
        current_row += 1
        current_row = add_excel_table(
            ws, current_row, headers, rows, f"TableSupport{index}"
        ) + 2
    format_sheet(ws)
    # Cette feuille contient plusieurs tableaux successifs : aucune ligne
    # ne doit rester figée lors du défilement.
    ws.freeze_panes = None
    # openpyxl conserve parfois « pane=bottomLeft » dans la sélection après
    # suppression du gel. Excel considère alors la vue comme incohérente et la
    # répare à l'ouverture. Nettoyer explicitement cette référence de volet.
    for selection in ws.sheet_view.selection:
        selection.pane = None


def write_standard_sheet(
    ws,
    rows: list[dict[str, object]],
    headers: list[str],
    table_name: str,
) -> None:
    add_excel_table(ws, 1, headers, rows, table_name)
    format_sheet(ws)


MODE_DEFINITIONS = {
    "matiere": {
        "label": "Économie matière",
        "sheet_suffix": "Matière",
        "file_suffix": "MATIERE",
    },
    "equilibre": {
        "label": "Équilibré",
        "sheet_suffix": "Équilibré",
        "file_suffix": "EQUILIBRE",
    },
    "fabrication": {
        "label": "Simplicité fabrication",
        "sheet_suffix": "Fabrication",
        "file_suffix": "FABRICATION",
    },
}


def optimization_mode_label(mode_key: str) -> str:
    return MODE_DEFINITIONS.get(mode_key, {}).get("label", mode_key)


def optimization_mode_sheet_suffix(mode_key: str) -> str:
    return MODE_DEFINITIONS.get(mode_key, {}).get("sheet_suffix", mode_key[:18])


def optimization_mode_file_suffix(mode_key: str) -> str:
    return MODE_DEFINITIONS.get(mode_key, {}).get("file_suffix", normalize_text(mode_key).upper())


def build_multi_mode_comparison_rows(
    mode_results: dict[str, dict[str, object]],
) -> list[dict[str, object]]:
    """
    Construit une comparaison homogène des KPI entre les modes sélectionnés.
    """
    if not mode_results:
        return []

    rendement_values = {
        key: parse_float(result.get("kpis", {}).get("Rendement matière (%)")) or 0
        for key, result in mode_results.items()
    }
    plan_values = {
        key: parse_float(result.get("kpis", {}).get("Plans de découpe distincts")) or 0
        for key, result in mode_results.items()
    }
    bar_values = {
        key: parse_float(result.get("kpis", {}).get("Barres optimisées")) or 0
        for key, result in mode_results.items()
    }
    length_values = {
        key: parse_float(result.get("kpis", {}).get("Longueur optimisée (m)")) or 0
        for key, result in mode_results.items()
    }

    max_rendement = max(rendement_values.values()) if rendement_values else 0
    min_plans = min(plan_values.values()) if plan_values else 0
    min_bars = min(bar_values.values()) if bar_values else 0
    min_length = min(length_values.values()) if length_values else 0

    rows: list[dict[str, object]] = []
    for mode_key, result in mode_results.items():
        kpis = result.get("kpis", {})
        strengths: list[str] = []
        if math.isclose(rendement_values[mode_key], max_rendement, abs_tol=1e-9):
            strengths.append("Meilleur rendement")
        if math.isclose(plan_values[mode_key], min_plans, abs_tol=1e-9):
            strengths.append("Moins de plans")
        if math.isclose(bar_values[mode_key], min_bars, abs_tol=1e-9):
            strengths.append("Moins de barres")
        if math.isclose(length_values[mode_key], min_length, abs_tol=1e-9):
            strengths.append("Moins de longueur achetée")

        rows.append(
            {
                "Mode": optimization_mode_label(mode_key),
                "Barres initiales": kpis.get("Barres initiales"),
                "Barres optimisées": kpis.get("Barres optimisées"),
                "Longueur initiale (m)": kpis.get("Longueur initiale (m)"),
                "Longueur optimisée (m)": kpis.get("Longueur optimisée (m)"),
                "Gain longueur (m)": kpis.get("Gain longueur (m)"),
                "Gain longueur (%)": kpis.get("Gain longueur (%)"),
                "Chute optimisée (m)": kpis.get("Chute optimisée (m)"),
                "Rendement matière (%)": kpis.get("Rendement matière (%)"),
                "Plans distincts": kpis.get("Plans de découpe distincts"),
                "Codes à chiffrer optimisés": kpis.get("Codes à chiffrer optimisés"),
                "Gain codes à chiffrer": kpis.get("Gain codes à chiffrer"),
                "Points forts": " ; ".join(strengths),
            }
        )
    return rows


def _write_cut_plan_sheet(
    ws,
    cut_summary: list[dict[str, object]],
    cut_detail: list[dict[str, object]],
    kpis: dict[str, float],
    table_suffix: str,
) -> None:
    ws.cell(1, 1, "KPI optimisation")
    ws.cell(1, 1).fill = HEADER_FILL
    ws.cell(1, 1).font = WHITE_FONT

    row_index = 2
    for key in [
        "Barres optimisées",
        "Longueur optimisée (m)",
        "Chute optimisée (m)",
        "Rendement matière (%)",
        "Plans de découpe distincts",
    ]:
        ws.cell(row_index, 1, key)
        value = kpis.get(key)
        ws.cell(row_index, 2, display_number(value))
        if key == "Rendement matière (%)" and isinstance(value, (int, float)):
            ws.cell(row_index, 2).number_format = "0.00"
        row_index += 1

    by_profile: list[dict[str, object]] = []
    for row in cut_summary:
        stock = (
            (parse_float(row.get("Taille des barres (m)")) or 0)
            * (parse_float(row.get("Nombre de barres")) or 0)
        )
        waste = parse_float(row.get("Chute totale (m)")) or 0
        used = parse_float(row.get("Longueur utile totale (m)")) or 0
        by_profile.append(
            {
                "Profil": row.get("Profil"),
                "Taille (m)": row.get("Taille des barres (m)"),
                "Barres": row.get("Nombre de barres"),
                "Longueur achetée (m)": stock,
                "Chute (m)": waste,
                "Rendement (%)": (used / stock * 100) if stock else 100,
            }
        )

    profile_headers = [
        "Profil",
        "Taille (m)",
        "Barres",
        "Longueur achetée (m)",
        "Chute (m)",
        "Rendement (%)",
    ]
    end_profile = add_excel_table(
        ws,
        row_index + 1,
        profile_headers,
        by_profile,
        f"KPIProfils{table_suffix}",
    )
    plan_start = end_profile + 2
    plan_headers = [
        "Profil",
        "Barre n°",
        "Taille de barre (m)",
        "Découpes (m)",
        "Longueur utilisée (m)",
        "Chute (m)",
        "Taux d'utilisation",
    ]
    add_excel_table(
        ws,
        plan_start,
        plan_headers,
        cut_detail,
        f"PlanDecoupe{table_suffix}",
    )
    format_sheet(ws)
    ws.freeze_panes = f"A{plan_start + 1}"


def create_output_workbook_multi(
    enriched_rows: list[dict[str, object]],
    total_rows: list[dict[str, object]],
    length_rows: list[dict[str, object]],
    piece_rows: list[dict[str, object]],
    initial_offer_rows: list[dict[str, object]],
    cable_rows: list[dict[str, object]],
    control_rows: list[dict[str, object]],
    mode_results: dict[str, dict[str, object]],
    selected_lengths: dict[str, list[float]],
    options: dict[str, ProfileOption],
    article_database: dict[str, dict[str, object]],
    source_csv: Path,
    source_pdf: Optional[Path],
    prioritize_exact: bool,
    replace_magnelis: bool,
    base_article_version: str,
    base_profiles_version: str,
) -> tuple[Workbook, list[dict[str, object]], dict[str, list[dict[str, object]]]]:
    """
    Génère un seul classeur contenant les résultats de tous les modes
    sélectionnés. Les besoins amont sont communs ; chaque mode conserve son
    propre Opti globale et son propre Plan de découpe.
    """
    wb = Workbook()
    wb.remove(wb.active)

    summary_headers = [
        "Type",
        "Code article",
        "Profil",
        "Libellé",
        "Category",
        "Quantité à chiffrer",
        "Unité",
        "Packaging",
        "Nombre de code à chiffrer",
    ]

    # Feuilles communes. Les contrôles sont générés mais masqués par défaut.
    ws_ctrl = wb.create_sheet(SHEET_CONTROLS)
    write_standard_sheet(
        ws_ctrl,
        control_rows,
        ["Contrôle", "Statut", "Détail"],
        "Controles",
    )
    for row in ws_ctrl.iter_rows(min_row=2):
        status = str(row[1].value or "")
        if status == "OK":
            row[1].fill = PatternFill("solid", fgColor="C6EFCE")
        elif status == "ERREUR":
            row[1].fill = PatternFill("solid", fgColor="FFC7CE")
        elif status == "ATTENTION":
            row[1].fill = PatternFill("solid", fgColor="FFEB9C")

    ws_support = wb.create_sheet(SHEET_BY_SUPPORT)
    write_by_support_sheet(ws_support, build_by_support(enriched_rows))

    ws_initial = wb.create_sheet(SHEET_INITIAL)
    write_standard_sheet(
        ws_initial,
        initial_offer_rows,
        summary_headers,
        "OffreInitiale",
    )

    ws_comparison = wb.create_sheet(SHEET_COMPARISON)
    comparison_rows = build_multi_mode_comparison_rows(mode_results)
    comparison_headers = [
        "Mode",
        "Barres initiales",
        "Barres optimisées",
        "Longueur initiale (m)",
        "Longueur optimisée (m)",
        "Gain longueur (m)",
        "Gain longueur (%)",
        "Chute optimisée (m)",
        "Rendement matière (%)",
        "Plans distincts",
        "Codes à chiffrer optimisés",
        "Gain codes à chiffrer",
        "Points forts",
    ]
    write_standard_sheet(
        ws_comparison,
        comparison_rows,
        comparison_headers,
        "ComparaisonModes",
    )

    optimized_by_mode: dict[str, list[dict[str, object]]] = {}
    piece_headers = ["Code article", "Libellé", "Quantité"]
    piece_output = [
        {header: row.get(header) for header in piece_headers}
        for row in piece_rows
        if not is_cable_tray_code(row.get("Code article"), article_database)
    ]

    # Une paire de feuilles visible par mode sélectionné pour l'étude Excel.
    for mode_key, result in mode_results.items():
        suffix = optimization_mode_sheet_suffix(mode_key)
        safe_suffix = re.sub(r"[^A-Za-z0-9]", "", normalize_text(suffix).title())[:14] or "Mode"
        cut_summary = result.get("cut_summary", [])
        cut_detail = result.get("cut_detail", [])
        kpis = result.get("kpis", {})

        optimized_rows = (
            enrich_summary_rows(cut_summary, piece_output, article_database)
            + cable_tray_offer_rows(cable_rows)
        )
        optimized_rows.sort(
            key=lambda row: (
                normalize_text(row.get("Category")),
                normalize_text(row.get("Code article")),
            )
        )
        optimized_by_mode[mode_key] = optimized_rows

        ws_summary = wb.create_sheet(f"Opti globale - {suffix}"[:31])
        write_standard_sheet(
            ws_summary,
            optimized_rows,
            summary_headers,
            f"OptiGlobale{safe_suffix}",
        )

        ws_plan = wb.create_sheet(f"Plan découpe - {suffix}"[:31])
        _write_cut_plan_sheet(
            ws_plan,
            cut_summary,
            cut_detail,
            kpis,
            safe_suffix,
        )

    ws_cable = wb.create_sheet(SHEET_CABLE_TRAY)
    cable_headers = [
        "Code article",
        "Libellé",
        "Category",
        "Métré nécessaire (m)",
        "Longueur commerciale (m)",
        "Nombre de barres",
        "Métré à chiffrer (m)",
        "Chute théorique (m)",
        "Unité",
        "Packaging",
    ]
    write_standard_sheet(
        ws_cable,
        cable_rows,
        cable_headers,
        "OptiCheminCables",
    )

    # Feuilles techniques masquées.
    ws_global = wb.create_sheet(SHEET_GLOBAL)
    write_global_sheet(ws_global, enriched_rows)

    total_headers = [
        "Code article",
        "Libellé",
        "Quantité",
        "Longueur utile",
        "Longueur proposée",
        "Chute calculée",
        "Métré nécessaire (m)",
        "Nombre de barres chemin de câbles",
    ]
    cable_map = {
        normalize_code(row.get("Code article")): row
        for row in cable_rows
    }
    total_output = []
    for row in total_rows:
        enriched = dict(row)
        cable = cable_map.get(normalize_code(row.get("Code article")))
        enriched["Métré nécessaire (m)"] = (
            cable.get("Métré nécessaire (m)") if cable else None
        )
        enriched["Nombre de barres chemin de câbles"] = (
            cable.get("Nombre de barres") if cable else None
        )
        total_output.append(enriched)

    ws_total = wb.create_sheet(SHEET_TOTAL)
    write_standard_sheet(
        ws_total,
        total_output,
        total_headers,
        "NomenclatureTotale",
    )

    ws_pieces = wb.create_sheet(SHEET_PIECES)
    write_standard_sheet(
        ws_pieces,
        piece_output,
        piece_headers,
        "OptiPieces",
    )

    for mode_key, result in mode_results.items():
        suffix = optimization_mode_sheet_suffix(mode_key)
        safe_suffix = re.sub(r"[^A-Za-z0-9]", "", normalize_text(suffix).title())[:14] or "Mode"
        ws_lengths = wb.create_sheet(f"Opti longueurs - {suffix}"[:31])
        length_headers = [
            "Code article",
            "Profil",
            "Catégorie",
            "Taille des barres (m)",
            "Nombre de barres",
            "Longueur utile totale (m)",
            "Chute totale (m)",
            "Taux d'utilisation",
        ]
        write_standard_sheet(
            ws_lengths,
            result.get("cut_summary", []),
            length_headers,
            f"OptiLongueurs{safe_suffix}",
        )
        ws_lengths.sheet_state = "hidden"

    ws_settings = wb.create_sheet(SHEET_SETTINGS)
    ws_settings.append(["Paramètre", "Valeur"])
    parameters = [
        ("CSV source", str(source_csv)),
        ("PDF source", str(source_pdf) if source_pdf else "Non utilisé"),
        (
            "Modes optimisation",
            ", ".join(optimization_mode_label(key) for key in mode_results),
        ),
        (
            "Priorité longueurs exactes",
            "Oui" if prioritize_exact else "Non",
        ),
        (
            "Remplacement Magnelis → EZ",
            "Oui" if replace_magnelis else "Non",
        ),
        ("Base article", base_article_version),
        ("Base profils", base_profiles_version),
    ]
    for key, value in parameters:
        ws_settings.append([key, value])
    for profile_id, selected in selected_lengths.items():
        ws_settings.append(
            [
                options[profile_id].profile_label,
                ", ".join(f"{value:g} m" for value in sorted(selected)),
            ]
        )

    # V1.1.3 : feuilles techniques masquées par défaut. Toutes les feuilles
    # commençant par "Opti" sont masquées sauf les "Opti globale".
    for ws in [ws_ctrl, ws_global, ws_total, ws_pieces, ws_settings, ws_cable]:
        ws.sheet_state = "hidden"
    for ws in wb.worksheets:
        if ws.title.startswith("Opti") and not ws.title.startswith("Opti globale"):
            ws.sheet_state = "hidden"

    # Mise en page d'impression appliquée à toutes les feuilles apparentes.
    for ws in wb.worksheets:
        if ws.sheet_state == "visible":
            configure_print_layout(ws)

    return wb, initial_offer_rows, optimized_by_mode



def create_output_workbook(
    enriched_rows: list[dict[str, object]], total_rows: list[dict[str, object]],
    length_rows: list[dict[str, object]], piece_rows: list[dict[str, object]],
    cut_summary: list[dict[str, object]], cut_detail: list[dict[str, object]],
    initial_offer_rows: list[dict[str, object]], cable_rows: list[dict[str, object]],
    control_rows: list[dict[str, object]], kpis: dict[str, float],
    selected_lengths: dict[str, list[float]], options: dict[str, ProfileOption],
    article_database: dict[str, dict[str, object]], source_csv: Path, source_pdf: Optional[Path],
    prioritize_exact: bool, replace_magnelis: bool, optimization_mode: str,
    base_article_version: str, base_profiles_version: str,
) -> tuple[Workbook, list[dict[str, object]], list[dict[str, object]]]:
    wb=Workbook(); wb.remove(wb.active)
    summary_headers=["Type","Code article","Profil","Libellé","Category","Quantité à chiffrer","Unité","Packaging","Nombre de code à chiffrer"]

    ws_ctrl=wb.create_sheet(SHEET_CONTROLS)
    write_standard_sheet(ws_ctrl, control_rows, ["Contrôle","Statut","Détail"], "Controles")
    for row in ws_ctrl.iter_rows(min_row=2):
        status=str(row[1].value or '')
        if status=='OK': row[1].fill=PatternFill('solid',fgColor='C6EFCE')
        elif status=='ERREUR': row[1].fill=PatternFill('solid',fgColor='FFC7CE')
        elif status=='ATTENTION': row[1].fill=PatternFill('solid',fgColor='FFEB9C')

    ws_support=wb.create_sheet(SHEET_BY_SUPPORT); write_by_support_sheet(ws_support,build_by_support(enriched_rows))
    ws_initial=wb.create_sheet(SHEET_INITIAL); write_standard_sheet(ws_initial,initial_offer_rows,summary_headers,'OffreInitiale')

    piece_headers=["Code article","Libellé","Quantité"]
    piece_output=[{h:r.get(h) for h in piece_headers} for r in piece_rows if not is_cable_tray_code(r.get('Code article'),article_database)]
    optimized_rows=enrich_summary_rows(cut_summary,piece_output,article_database)
    optimized_rows.extend(cable_tray_offer_rows(cable_rows)); optimized_rows.sort(key=lambda r:(normalize_text(r.get('Category')),normalize_text(r.get('Code article'))))
    ws_summary=wb.create_sheet(SHEET_SUMMARY); write_standard_sheet(ws_summary,optimized_rows,summary_headers,'OptiGlobale')

    ws_cable=wb.create_sheet(SHEET_CABLE_TRAY)
    cable_headers=["Code article","Libellé","Category","Métré nécessaire (m)","Longueur commerciale (m)","Nombre de barres","Métré à chiffrer (m)","Chute théorique (m)","Unité","Packaging"]
    write_standard_sheet(ws_cable,cable_rows,cable_headers,'OptiCheminCables')

    ws_comp=wb.create_sheet(SHEET_COMPARISON)
    comp_rows=[{"Indicateur":k,"Valeur":v} for k,v in kpis.items()]
    write_standard_sheet(ws_comp,comp_rows,["Indicateur","Valeur"],'SyntheseOptimisation')

    ws_plan=wb.create_sheet(SHEET_CUT_PLAN)
    # KPI block above plan
    ws_plan.cell(1,1,'KPI optimisation'); ws_plan.cell(1,1).fill=HEADER_FILL; ws_plan.cell(1,1).font=WHITE_FONT
    rr=2
    for k in ["Barres optimisées","Longueur optimisée (m)","Chute optimisée (m)","Rendement matière (%)","Plans de découpe distincts"]:
        ws_plan.cell(rr,1,k); ws_plan.cell(rr,2,display_number(kpis.get(k))); rr+=1
    # KPI par profil
    by_profile=[]
    for row in cut_summary:
        stock=(parse_float(row.get('Taille des barres (m)')) or 0)*(parse_float(row.get('Nombre de barres')) or 0)
        waste=parse_float(row.get('Chute totale (m)')) or 0
        used=parse_float(row.get('Longueur utile totale (m)')) or 0
        by_profile.append({'Profil':row.get('Profil'),'Taille (m)':row.get('Taille des barres (m)'),'Barres':row.get('Nombre de barres'),'Longueur achetée (m)':stock,'Chute (m)':waste,'Rendement (%)':(used/stock*100) if stock else 100})
    profile_headers=['Profil','Taille (m)','Barres','Longueur achetée (m)','Chute (m)','Rendement (%)']
    end_profile=add_excel_table(ws_plan,rr+1,profile_headers,by_profile,'KPIProfils')
    plan_start=end_profile+2
    plan_headers=["Profil","Barre n°","Taille de barre (m)","Découpes (m)","Longueur utilisée (m)","Chute (m)","Taux d'utilisation"]
    add_excel_table(ws_plan,plan_start,plan_headers,cut_detail,'PlanDecoupe'); format_sheet(ws_plan); ws_plan.freeze_panes=f'A{plan_start+1}'

    # hidden technical sheets
    ws_global=wb.create_sheet(SHEET_GLOBAL); write_global_sheet(ws_global,enriched_rows)
    total_headers=["Code article","Libellé","Quantité","Longueur utile","Longueur proposée","Chute calculée","Métré nécessaire (m)","Nombre de barres chemin de câbles"]
    # enrich total rows with cable aggregate
    cmap={normalize_code(r.get('Code article')):r for r in cable_rows}
    total_out=[]
    for r in total_rows:
        x=dict(r); c=cmap.get(normalize_code(r.get('Code article')))
        x['Métré nécessaire (m)']=c.get('Métré nécessaire (m)') if c else None
        x['Nombre de barres chemin de câbles']=c.get('Nombre de barres') if c else None
        total_out.append(x)
    ws_total=wb.create_sheet(SHEET_TOTAL); write_standard_sheet(ws_total,total_out,total_headers,'NomenclatureTotale')
    length_headers=["Code article","Profil","Catégorie","Taille des barres (m)","Nombre de barres","Longueur utile totale (m)","Chute totale (m)","Taux d'utilisation"]
    ws_lengths=wb.create_sheet(SHEET_LENGTHS); write_standard_sheet(ws_lengths,cut_summary,length_headers,'OptiLongueurs')
    ws_pieces=wb.create_sheet(SHEET_PIECES); write_standard_sheet(ws_pieces,piece_output,piece_headers,'OptiPieces')
    ws_settings=wb.create_sheet(SHEET_SETTINGS); ws_settings.append(['Paramètre','Valeur'])
    for k,v in [
        ('CSV source',str(source_csv)),('PDF source',str(source_pdf) if source_pdf else 'Non utilisé'),
        ('Mode optimisation',optimization_mode),('Priorité longueurs exactes','Oui' if prioritize_exact else 'Non'),
        ('Remplacement Magnelis → EZ','Oui' if replace_magnelis else 'Non'),('Base article',base_article_version),('Base profils',base_profiles_version)
    ]: ws_settings.append([k,v])
    for pid,selected in selected_lengths.items(): ws_settings.append([options[pid].profile_label,', '.join(f'{v:g} m' for v in sorted(selected))])
    for ws in [ws_global,ws_total,ws_lengths,ws_pieces,ws_settings]: ws.sheet_state='hidden'
    return wb,initial_offer_rows,optimized_rows

